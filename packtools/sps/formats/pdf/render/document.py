"""Monta o DOCX: seções (capa, corpo em colunas, anexos), estilos e numeração de páginas."""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from packtools.sps.formats.pdf.layout import page as page_layout
from packtools.sps.formats.pdf.render import front as front_render
from packtools.sps.formats.pdf.render.blocks import render_blocks
from packtools.sps.formats.pdf.render.styles import copy_styles


class State:
    """Documento em construção e o que os handlers precisam saber sobre ele."""

    def __init__(self, doc, page, assets):
        """doc: Document em construção; page: atributos de página; assets: resolvedor de imagens."""
        self.doc = doc
        self.page = page
        self.assets = assets
        self.multi_column = page["column_count"] > 1
        self.full_width = False
        # True na capa e nos anexos: seções já em coluna única, onde abrir
        # outra seção de coluna única para tabela/figura larga não faz sentido
        self.locked = False
        # nomes das marcas já inseridas (alvos de <xref>); a mesma marca não se repete
        self.bookmarks = set()

    def body_section(self):
        """Segunda seção (corpo); criada com o tipo de início configurado."""
        # Herdado do old_pdf (#1295/#1296): o padrão (default_layout.json,
        # body_start_type="continuous") é quebra de seção contínua, para o corpo
        # começar já na página 1, logo abaixo dos resumos, e não numa página nova.
        while len(self.doc.sections) < 2:
            self.doc.add_section(self.page["start_type"])
        return self.doc.sections[1]

    def set_full_width(self, full_width):
        """Abre/fecha uma seção contínua de coluna única para tabela/figura largas."""
        # Herdado do old_pdf: no DOCX um elemento não pode ser mais largo que a
        # coluna; para uma tabela/figura ocupar a largura da página, abre-se uma
        # seção contínua de 1 coluna e, depois dela, volta-se às colunas do corpo.
        if self.locked or not self.multi_column or full_width == self.full_width:
            return
        section = self.doc.add_section(WD_SECTION.CONTINUOUS)
        if full_width:
            set_columns(section, 1, 0)
        else:
            set_columns(section, self.page["column_count"], self.page["column_spacing"])
        self.full_width = full_width

    def add_inline_image(self, run, href):
        """
        Insere imagem inline no run (largura natural limitada à coluna); False se não houver arquivo.

        Exceção à regra de que as medidas chegam prontas do layout: uma
        imagem inline é um run dentro de um parágrafo (ou de uma célula),
        não um bloco, e ``layout.placement.place`` só percorre blocos. A
        regra continua no layout (``layout.page.fit_width``); aqui ela só é
        chamada na hora de escrever.
        """
        path = self.assets.path(href)
        if not path:
            return False
        width = page_layout.fit_width(
            self.assets.natural_width(href), page_layout.column_width(self.page)
        )
        try:
            run.add_picture(path, width=width)
        except Exception:
            return False
        return True


def new_document(base_layout=None):
    """Document vazio com os estilos do layout base (.docx), quando informado."""
    doc = Document()
    if base_layout:
        copy_styles(Document(base_layout), doc)
    return doc


def render(article, page, assets, base_layout=None, logo_spec=None):
    """
    Escreve o artigo extraído: capa (coluna única), cabeçalhos e rodapés,
    corpo em colunas, anexos em seção própria e, por fim, atributos de
    página e numeração inicial em todas as seções.
    """
    state = State(new_document(base_layout), page, assets)
    meta = article["meta"]

    state.locked = True  # capa em coluna única
    front_render.render_front(state, meta, article["front"], logo_spec)
    state.locked = False
    front_render.render_headers_footers(state, meta)

    body = state.body_section()
    set_columns(
        body, page["column_count"], page["column_spacing"] if state.multi_column else 0
    )
    render_blocks(state, article["main"])

    # Herdado do old_pdf (regra editorial): apêndices e material suplementar
    # começam em página nova (add_section sem tipo = nova página), em coluna
    # única e com rodapé próprio, sem número de página: o material
    # suplementar tem paginação independente da do artigo.
    if article["annex"]:
        annex = state.doc.add_section()
        front_render.annex_footer(state, annex, meta)
        set_columns(annex, 1, page["column_spacing"])
        state.locked = True
        render_blocks(state, article["annex"])

    _apply_page(state.doc, page)
    # Herdado do old_pdf (#1295, #1296, #1302): a numeração começa em fpage
    # (página editorial do artigo na revista). O LibreOffice só respeita o
    # reinício se ele estiver em TODAS as seções: só na do corpo é ignorado
    # (sai a contagem física), e só na primeira faz a quebra contínua virar
    # quebra de página. É fpage (e não fpage + 1) porque o campo PAGE já
    # incrementa sozinho. Sem fpage (publicação contínua, com elocation-id),
    # as folhas são numeradas na ordem natural: a primeira é 1, não 0.
    try:
        start = int(meta["fpage"])
    except ValueError:
        start = 1
    for section in state.doc.sections:
        _set_start_page(section, start)
    return state.doc


def set_columns(section, count, spacing_twips):
    """Define número de colunas e espaçamento (twips) da seção (w:cols)."""
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(spacing_twips))


def _apply_page(doc, page):
    """Aplica margens, tamanho e orientação a todas as seções; 1ª página com cabeçalho/rodapé próprios."""
    doc.sections[0].different_first_page_header_footer = page[
        "different_first_page_header_footer"
    ]
    for section in doc.sections:
        for attr in (
            "top_margin",
            "left_margin",
            "right_margin",
            "bottom_margin",
            "header_distance",
            "footer_distance",
            "gutter",
            "orientation",
            "page_height",
            "page_width",
        ):
            setattr(section, attr, page[attr])


def _set_start_page(section, number):
    """Número inicial da paginação da seção (w:pgNumType/@w:start)."""
    pg_num_type = section._sectPr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        section._sectPr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), str(number))
