"""Blocos -> DOCX: um handler por tipo e troca de colunas para elementos de largura total."""

import logging

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from packtools.sps.formats.pdf.config.page import SINGLE_COLUMN
from packtools.sps.formats.pdf.extract.runs import has_content
from packtools.sps.formats.pdf.render import styles
from packtools.sps.formats.pdf.render.inline import add_runs
from packtools.sps.formats.pdf.render.table import add_caption, add_table

logger = logging.getLogger(__name__)

LIST_INDENT = Cm(0.63)
# Estilo usado quando o layout base não tem o do papel (ex. sem layout.docx).
ROLE_FALLBACKS = {"article_title": "Heading 1", "abstract_title": "Heading 2"}
# Herdado do old_pdf (#1322): SCL Paragraph Keywords não tem espaço depois, e
# o título de seção seguinte (sem espaço antes) ficava colado nas palavras-chave.
# 5.65pt = espaço depois do SCL Paragraph, mantendo o ritmo do texto.
KEYWORDS_SPACE_AFTER = Pt(5.65)


def render_blocks(state, blocks):
    """Escreve os blocos em ordem, abrindo/fechando seção de coluna única conforme block["layout"]."""
    for block in blocks:
        state.set_full_width(block.get("layout") == SINGLE_COLUMN)
        RENDERERS[block["type"]](state, block)
    state.set_full_width(False)


def _paragraph(state, role, fallback="body"):
    """Novo parágrafo com o estilo do papel (ou o de fallback, se o layout não tiver o primeiro)."""
    paragraph = state.doc.add_paragraph()
    styles.apply(paragraph, state.doc, role) or styles.apply(
        paragraph, state.doc, fallback
    )
    return paragraph


def heading(state, block):
    """Título de seção: nível 2 -> "SCL Section Title"; 3 ou mais -> "SCL Subsection Title"."""
    # Nível 4 ou mais também usa 'SCL Subsection Title' (o old_pdf usava
    # 'SCL Paragraph', e o título ficava igual a um parágrafo de texto).
    role = "heading2" if block["level"] <= 2 else "heading3"
    paragraph = _paragraph(state, role, f"Heading {min(block['level'], 9)}")
    add_runs(paragraph, block["runs"], state)


def paragraph(state, block):
    """
    Parágrafo de texto (do corpo ou da capa): estilo do papel, estilo de
    caractere (block["char_role"], usado na capa), recuo de lista, entrelinha
    simples nas referências e espaço depois das palavras-chave.
    """
    role = block["role"]
    para = _paragraph(state, role, ROLE_FALLBACKS.get(role, "body"))
    if block.get("indent"):
        para.paragraph_format.left_indent = LIST_INDENT * block["indent"]
    if role == "reference":
        # SCL Paragraph Reference não define entrelinha: sem isso a referência longa fica espaçada como o corpo
        para.paragraph_format.line_spacing = 1.0
    elif role == "keywords":
        para.paragraph_format.space_after = KEYWORDS_SPACE_AFTER
    add_runs(para, block["runs"], state, char_role=block.get("char_role"))


def formula(state, block):
    """Fórmula em bloco: parágrafo "SCL Formula" com o OMML e o rótulo."""
    add_runs(_paragraph(state, "formula"), block["runs"], state)


def figure(state, block):
    """
    Figura: imagem no tamanho anotado por layout.placement.place
    (block["width"] e block["height"], em EMU) e legenda depois (ou antes,
    para tabela publicada como imagem). Sem imagem legível, imprime o texto
    alternativo.
    """
    # alvo de <xref ref-type="fig">: a marca vai no parágrafo da imagem
    mark = [{"bookmark": block["anchor"]}] if block.get("anchor") else []
    if block["caption_first"] and has_content(block["caption"]):
        add_caption(state, mark + block["caption"], block["caption_role"])
        mark = []
    picture = _picture(state, block["href"], block["width"], block["height"])
    if picture is not None:
        add_runs(picture, mark, state)
    else:
        add_runs(
            _paragraph(state, "body"),
            mark + [{"text": block["alt"] or block["href"] or "[Figure]"}],
            state,
        )
    if not block["caption_first"] and has_content(block["caption"]):
        add_caption(state, block["caption"], block["caption_role"])


def _picture(state, href, width, height):
    """
    Imagem centrada no tamanho dado (EMU); devolve o parágrafo, ou None sem imagem.

    O tamanho já vem decidido pelo layout (ver layout.page.fit_image):
    aqui não se reduz nem se amplia nada. height None (o layout não
    conseguiu ler a imagem) deixa o python-docx calcular a altura pela
    proporção que ele lê do arquivo.
    """
    path = state.assets.path(href)
    if not path:
        return None
    para = state.doc.add_paragraph()
    try:
        # Herdado do old_pdf: largura sempre explícita. Sem ela, o python-docx lê o
        # DPI por conta própria (72 quando falta o metadado) e a imagem sai em
        # tamanho diferente do calculado no layout.
        picture = para.add_run().add_picture(
            path, width=int(width), height=None if height is None else int(height)
        )
    except Exception as exc:
        logger.warning(
            "imagem %s não inserida no DOCX; sai o texto alternativo (%s: %s)",
            href,
            type(exc).__name__,
            exc,
        )
        para._p.getparent().remove(para._p)
        return None
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # sem dist*, o LibreOffice aplica 0,32 cm à esquerda e descentraliza a figura
    for side in ("distT", "distB", "distL", "distR"):
        picture._inline.set(side, "0")
    return para


RENDERERS = {
    "heading": heading,
    "paragraph": paragraph,
    "formula": formula,
    "figure": figure,
    "table": add_table,
}
