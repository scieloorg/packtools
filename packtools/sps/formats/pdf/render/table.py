"""
Bloco 'table' (um por <table>) -> legenda, tabela DOCX e notas de rodapé da tabela.

Não decide medidas: a largura da tabela e a de cada coluna chegam prontas
em ``block["width"]`` e ``block["column_widths"]`` (EMU), anotadas por
``layout.placement.place``. Aqui só se escreve o XML que as aplica.
"""

from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from packtools.sps.formats.pdf.config.page import SINGLE_COLUMN
from packtools.sps.formats.pdf.extract.runs import has_content
from packtools.sps.formats.pdf.render import styles
from packtools.sps.formats.pdf.render.inline import add_runs

# Herdado do old_pdf: tabelas em 7pt (células e notas), para caber na coluna.
CELL_FONT_SIZE = 7
FOOT_FONT_SIZE = 7


def add_caption(state, runs, role):
    """Legenda justificada; a de tabela tem entrelinha simples e fica presa à tabela."""
    paragraph = state.doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if role == "table_caption":
        # Herdado do old_pdf: SCL Table Heading não define entrelinha; sem isso a
        # legenda longa fica espaçada como o texto do corpo.
        paragraph.paragraph_format.line_spacing = 1.0
    if styles.apply(paragraph, state.doc, role) is not None and role == "table_caption":
        # a legenda não fica sozinha no fim da página, separada da tabela
        paragraph.paragraph_format.keep_with_next = True
    # só o rótulo (já marcado bold na extração) fica em negrito
    add_runs(
        paragraph,
        [dict(r, bold=r.get("bold", False)) if "text" in r else r for r in runs],
        state,
    )
    return paragraph


def add_table(state, block):
    """
    Tabela: legenda acima (só no 1º painel), a tabela DOCX (largura e
    colunas conforme anotadas por layout.placement.place) e as notas abaixo
    (só no último painel), em itálico 7pt. Um <table-wrap> com várias <table> chega como
    blocos consecutivos, um por painel (ver extract.table_wrap).
    """
    # alvo de <xref ref-type="table">: a marca vai na legenda ou, sem ela, na 1ª célula
    mark = [{"bookmark": block["anchor"]}] if block.get("anchor") else []
    if has_content(block["caption"]):
        add_caption(state, mark + block["caption"], "table_caption")
    elif mark and block["cells"]:
        first = block["cells"][0]
        first["runs"] = mark + first["runs"]
    single = block["layout"] == SINGLE_COLUMN
    if block["ncols"]:
        # Flutua só quando interrompe o corpo em colunas; em seção de coluna única
        # (anexos) fica no fluxo, podendo quebrar entre páginas. Fica no render
        # (e não no layout) porque depende do estado da seção em construção.
        _add_grid(
            state,
            block,
            single,
            floating=single and state.multi_column and not state.locked,
        )
    for i, runs in enumerate(block["foot"]):
        paragraph = state.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        # Herdado do old_pdf: notas em itálico 7pt, separadas da tabela por 6pt.
        paragraph.paragraph_format.space_before = Pt(6) if i == 0 else Pt(0)
        add_runs(
            paragraph,
            [dict(r, italic=True) if "text" in r else r for r in runs],
            state,
            size=FOOT_FONT_SIZE,
        )


def _add_grid(state, grid, single, floating):
    """
    Uma grade (bloco table: nrows, ncols, cells, width, column_widths) ->
    tabela DOCX: largura fixa, células mescladas pelos spans,
    conteúdo 7pt (cabeçalho em negrito), estilo 'Table Grid' com bordas
    superior/inferior; floating=True a torna flutuante na largura toda.
    """
    table = state.doc.add_table(rows=max(1, grid["nrows"]), cols=grid["ncols"])
    table.autofit = False
    tbl_pr = table._element.tblPr
    # w:tblW em twips (dxa): 635 EMU por twip
    _replace(tbl_pr, "w:tblW", **{"w:w": str(int(grid["width"] / 635)), "w:type": "dxa"})
    _replace(tbl_pr, "w:tblCellSpacing", **{"w:w": "0", "w:type": "dxa"})
    _set_cell_margins(tbl_pr, "w:tblCellMar")
    _replace(tbl_pr, "w:tblLayout", **{"w:type": "fixed"})
    _remove(tbl_pr, "w:tblInd")
    _remove(tbl_pr, "w:jc")
    if not single:
        _replace(tbl_pr, "w:jc", **{"w:val": "center"})

    widths = grid["column_widths"]
    for column, column_width_ in zip(table.columns, widths):
        column.width = column_width_
    # O LibreOffice usa a largura de cada célula (w:tcW), não só a da coluna.
    # Ajustada antes das mesclagens: o merge do python-docx soma as larguras.
    for row in table.rows:
        for docx_cell, column_width_ in zip(row.cells, widths):
            docx_cell.width = column_width_

    for cell in grid["cells"]:
        if cell["rowspan"] > 1 or cell["colspan"] > 1:
            last = table.cell(
                cell["row"] + cell["rowspan"] - 1,
                min(cell["col"] + cell["colspan"], grid["ncols"]) - 1,
            )
            table.cell(cell["row"], cell["col"]).merge(last)
    for cell in grid["cells"]:
        width = sum(widths[cell["col"] : cell["col"] + cell["colspan"]])
        _fill_cell(state, table.cell(cell["row"], cell["col"]), cell, width)

    _set_row_heights(table, grid)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if styles.get(state.doc, "Table Grid") is not None:
        table.style = "Table Grid"
    # Herdado do old_pdf: linhas horizontais em cima e embaixo da tabela.
    _set_outer_borders(tbl_pr)
    if floating:
        _float_full_width(table)


def _fill_cell(state, docx_cell, cell, width):
    """Escreve o conteúdo da célula: 7pt, 1ª coluna à esquerda e demais centradas, sem espaçamento nem margens; imagem limitada à largura (EMU) da célula."""
    paragraph = docx_cell.paragraphs[0]
    add_runs(
        paragraph,
        cell["runs"],
        state,
        size=CELL_FONT_SIZE,
        bold=True if cell["header"] else None,
        image_width=width,
    )
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER if cell["col"] >= 1 else WD_ALIGN_PARAGRAPH.LEFT
    )
    fmt = paragraph.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    fmt.keep_together = True
    fmt.keep_with_next = fmt.page_break_before = fmt.widow_control = False
    _set_cell_margins(docx_cell._element.get_or_add_tcPr(), "w:tcMar")


def _set_row_heights(table, grid):
    """
    Altura mínima livre; linha só coberta por mesclagem vertical, sem conteúdo
    próprio, fica com 1pt. Linha só com "-"/"—" é dado e continua visível
    (o old_pdf também a reduzia a 1pt).
    """
    starts = {c["row"] for c in grid["cells"] if has_content(c["runs"])}
    covered = {c["row"] + i for c in grid["cells"] for i in range(1, c["rowspan"])}
    for index, row in enumerate(table.rows):
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.height = Pt(0)
        for docx_cell in row.cells:
            docx_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        if index in covered and index not in starts:
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            row.height = Pt(1)


def _float_full_width(table):
    """Tabela flutuante (sem sobreposição), centrada entre as margens, linhas sem quebra entre páginas."""
    tbl_pr = table._element.tblPr
    _replace(tbl_pr, "w:tblOverlap", **{"w:val": "never"})
    _replace(
        tbl_pr,
        "w:tblpPr",
        **{
            "w:horzAnchor": "margin",
            "w:vertAnchor": "paragraph",
            "w:tblpXSpec": "center",
            "w:leftFromText": "0",
            "w:rightFromText": "0",
            "w:topFromText": "0",
            "w:bottomFromText": "0",
        },
    )
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))


def _set_outer_borders(tbl_pr, color="000000", size=8):
    """Bordas superior e inferior da tabela (w:tblBorders)."""
    borders = _replace(tbl_pr, "w:tblBorders")
    for side in ("top", "bottom"):
        border = OxmlElement(f"w:{side}")
        for key, value in (
            ("w:val", "single"),
            ("w:sz", str(size)),
            ("w:color", color),
            ("w:space", "0"),
        ):
            border.set(qn(key), value)
        borders.append(border)


def _set_cell_margins(parent, tag):
    """Margens zeradas em todos os lados (w:tblCellMar da tabela ou w:tcMar da célula)."""
    margins = _replace(parent, tag)
    for side in ("top", "left", "bottom", "right"):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), "0")
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)


def _remove(parent, tag):
    """Remove o filho tag de parent, se existir."""
    existing = parent.find(qn(tag))
    if existing is not None:
        parent.remove(existing)


def _replace(parent, tag, **attrs):
    """Substitui o filho tag de parent por um novo com os atributos dados; devolve o novo elemento."""
    _remove(parent, tag)
    element = OxmlElement(tag)
    for key, value in attrs.items():
        element.set(qn(key), value)
    parent.append(element)
    return element
