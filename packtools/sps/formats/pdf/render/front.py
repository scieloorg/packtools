"""
Primeira página (cabeçalho com título/DOI/logo e os blocos da capa) e cabeçalhos/rodapés.

As larguras das colunas do cabeçalho e o tamanho do logo vêm de
``layout.logo`` (title_doi_columns, logo_header); aqui só se escrevem as
tabelas, o texto e a imagem.
"""

import logging

from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt
from docx.text.paragraph import Paragraph
from packtools.sps.formats.pdf.layout import logo as logo_layout
from packtools.sps.formats.pdf.layout import page as page_layout
from packtools.sps.formats.pdf.render import logo as logo_render
from packtools.sps.formats.pdf.render import styles
from packtools.sps.formats.pdf.render.blocks import render_blocks

logger = logging.getLogger(__name__)


def render_front(state, meta, blocks, logo_spec=None):
    """
    Primeira página: cabeçalho com título do periódico e DOI (ou com logo)
    e os blocos da capa.

    Parâmetros
    ----------
    state : render.document.State
        Documento em construção.
    meta : dict
        Metadados (ver extract.article_meta.extract_meta).
    blocks : list of dict
        Blocos da capa: categoria, título, autores, afiliações, notas de
        autor e resumos (ver extract.article_meta.front_blocks), escritos
        como os do corpo.
    logo_spec : layout.logo.LogoSpec, opcional
        Logo do periódico; sem ele (ou ilegível), título em texto.
    """
    doc = state.doc
    header = doc.sections[0].first_page_header
    if not (logo_spec and _logo_header(state, header, logo_spec, meta)):
        title_cell, doi_cell = _header_table(
            state,
            header,
            logo_layout.title_doi_columns(page_layout.content_width(state.page)),
            v_center=False,
        )
        _add_journal_title(
            title_cell.paragraphs[0], meta["journal_title"], title_cell.width, doc
        )
        _add_doi(doi_cell.paragraphs[0], meta["doi"], doc, WD_ALIGN_PARAGRAPH.RIGHT)
    _zero_cell_spacing(header)
    render_blocks(state, blocks)


def render_headers_footers(state, meta):
    """
    Cabeçalho das demais páginas (periódico | título do artigo), rodapé da
    1ª página (CITE AS + página | volume) e rodapé das demais páginas.
    meta: metadados (ver extract.article_meta.extract_meta).
    """
    doc = state.doc

    # Demais páginas: título do periódico | título do artigo.
    # Herdado do old_pdf: fica no cabeçalho padrão da 1ª seção (e não num
    # cabeçalho próprio da seção do corpo) porque o LibreOffice não mostra o
    # cabeçalho de uma seção com quebra contínua desvinculada da anterior, e
    # o corpo precisa da quebra contínua para começar na página 1. É uma
    # tabela de 2 colunas, e não texto com tabulação, porque tabulação só
    # posiciona o início do texto: um título de artigo longo quebraria para
    # baixo do título do periódico. Aqui o título do periódico vai numa linha
    # só e em fonte pequena; o tratamento em 2 linhas da 1ª página o faria
    # ocupar 3 linhas nesta coluna mais estreita.
    journal_cell, title_cell = _header_table(
        state,
        doc.sections[0].header,
        page_layout.split_width(page_layout.content_width(state.page), 0.5),
        v_center=False,
    )
    for cell, text, alignment in (
        (journal_cell, meta["journal_title"], None),
        (title_cell, meta["title_text"], WD_ALIGN_PARAGRAPH.RIGHT),
    ):
        paragraph = cell.paragraphs[0]
        styles.apply(paragraph, doc, "header")
        if alignment is not None:
            paragraph.alignment = alignment
        _add_text(paragraph, text, doc, "header_char")

    # Rodapé da 1ª página: CITE AS + página | volume.
    # Herdado do old_pdf (#1349): a nota "como citar" do XML já é a citação
    # completa e sai como está (acrescentar periódico/volume duplicava o que
    # ela já diz), só com ponto final se não tiver pontuação no fim. Sem
    # nota, a citação é montada pelos metadados (extract.citation.build_citation);
    # periódico + volume: local é o último recurso, sem o ": " solto quando
    # não há volume.
    first_footer = doc.sections[0].first_page_footer
    cite = (
        first_footer.paragraphs[0]
        if first_footer.paragraphs
        else first_footer.add_paragraph()
    )
    styles.apply(cite, doc, "cite_as")
    _add_text(cite, "CITE AS: ", doc, "cite_as_label_char")
    if meta["cite_as"]:
        text = meta["cite_as"]
        _add_text(
            cite,
            text + ("" if text.endswith((".", "!", "?")) else "."),
            doc,
            "cite_as_char",
        )
    else:
        _add_text(cite, f"{meta['journal_title']} ", doc, "cite_as_journal_char")
        part = (
            f"{meta['volume']}: {meta['location_label']}"
            if meta["volume"]
            else meta["location_label"]
        )
        _add_text(cite, f"{part}.", doc, "cite_as_char")
    _page_footer(state, first_footer.add_paragraph(), meta)

    # demais páginas
    body_footer = state.body_section().footer
    _page_footer(state, body_footer.paragraphs[0], meta)


def annex_footer(state, section, meta):
    """Anexos têm paginação própria: só volume/ano, sem número de página."""
    # Herdado do old_pdf (regra editorial): por isso também não há o " | " inicial.
    section.footer.is_linked_to_previous = False
    paragraph = section.footer.add_paragraph()
    styles.apply(paragraph, state.doc, "footer")
    paragraph.add_run(vol_issue_year(meta))


def vol_issue_year(footer):
    """
    'VOL. v (n) ano: local', sem as partes ausentes (e sem ':' quando falta o local).
    footer: dict com volume, issue, year e location_label (os metadados).
    """
    parts = [f"VOL. {footer['volume']}"] if footer["volume"] else []
    if footer["issue"]:
        parts.append(f"({footer['issue']})")
    if footer["year"]:
        parts.append(footer["year"])
    text = " ".join(parts)
    # publicação contínua sem elocation-id não tem local: evita "2023:" com ':' solto
    if footer["location_label"]:
        text = (
            f"{text}: {footer['location_label']}" if text else footer["location_label"]
        )
    return text


def journal_title_two_lines(title):
    """Primeira palavra numa linha, o resto na segunda (cabeçalho da 1ª página)."""
    # Herdado do old_pdf: limita o título do periódico, em fonte grande, a 2 linhas.
    first, _, rest = title.partition(" ")
    return f"{first}\n{rest}" if rest else title


# Largura estimada do título do periódico, que sai em caixa alta: média, em
# em, das maiúsculas da Noto Serif (a fonte do título no layout.docx) por
# classe de caractere. Medida com as fontes reais nos 315 periódicos dos
# lotes de teste, erra até 4,4% para menos; a margem de 2% evita que uma
# linha estimada como cabendo quebre de novo e o título ocupe 3 linhas.
_NARROW_CHARS = set(".,:;/-()'IJÍ")
_MEDIUM_CHARS = set("SFZPTCELYÇÊÉ")
_WIDE_CHARS = set("MW")
_WIDTH_MARGIN = 1.02
JOURNAL_TITLE_SIZE = 14
JOURNAL_TITLE_MIN_SIZE = 10


def _caps_width_em(text):
    """Largura estimada (em) do texto em caixa alta, com a margem de segurança."""
    width = 0.0
    for char in text.upper():
        if char == " ":
            width += 0.26
        elif char in _NARROW_CHARS:
            width += 0.33
        elif char in _MEDIUM_CHARS:
            width += 0.61
        elif char in _WIDE_CHARS:
            width += 0.95
        else:
            width += 0.72
    return width * _WIDTH_MARGIN


def fit_journal_title(title, width, size=JOURNAL_TITLE_SIZE):
    """
    Título do periódico em até 2 linhas na coluna do cabeçalho da 1ª página.

    Parâmetros
    ----------
    title : str
        Título do periódico.
    width : int
        Largura da coluna do título (EMU).
    size : float
        Tamanho da fonte do título (pt), o do estilo.

    Retorna
    -------
    tuple of (str, float)
        O texto, com a quebra de linha, e o tamanho da fonte. Na ordem: a
        quebra de journal_title_two_lines, se couber; senão, a que divide
        o título de forma mais equilibrada (a de menor linha mais longa);
        senão, a equilibrada com a fonte reduzida de 1 em 1 pt até
        JOURNAL_TITLE_MIN_SIZE, onde fica mesmo sem caber.
    """
    words = title.split()
    if len(words) < 2:
        return title, size
    splits = [
        (" ".join(words[:i]), " ".join(words[i:])) for i in range(1, len(words))
    ]

    def widest(split):
        return max(_caps_width_em(line) for line in split)

    balanced = min(splits, key=widest)
    current_size = size
    while True:
        limit = Emu(width).pt / current_size
        if current_size == size and widest(splits[0]) <= limit:
            return "\n".join(splits[0]), size
        if widest(balanced) <= limit or current_size - 1 < JOURNAL_TITLE_MIN_SIZE:
            return "\n".join(balanced), current_size
        current_size -= 1


# -----------------
# Partes da primeira página
# -----------------


def _add_text(paragraph, text, doc, char_role):
    """Acrescenta um run de texto com o estilo de caractere do papel (se existir)."""
    run = paragraph.add_run(text)
    style = styles.get(doc, char_role)
    if style is not None:
        run.style = style
    return run


def _add_journal_title(paragraph, title, width, doc):
    """
    Título do periódico em até 2 linhas na coluna de largura width (EMU),
    no tamanho do estilo ou reduzido (ver fit_journal_title).
    """
    style = styles.get(doc, "journal_title_char")
    size = JOURNAL_TITLE_SIZE
    if style is not None and style.font.size:
        size = style.font.size.pt
    text, fitted = fit_journal_title(title, width, size)
    run = _add_text(paragraph, text, doc, "journal_title_char")
    if fitted != size:
        run.font.size = Pt(fitted)


def _add_doi(paragraph, doi, doc, alignment):
    """DOI como URL no parágrafo, com o alinhamento dado."""
    paragraph.alignment = alignment
    if doi:
        _add_text(paragraph, f"http://dx.doi.org/{doi}", doc, "header_char")


def _page_footer(state, paragraph, meta):
    """Linha "PÁGINA | VOL. v (n) ano: local" com campo PAGE dinâmico."""
    # Herdado do old_pdf: campo PAGE dinâmico (e não fpage como texto fixo)
    # nos dois rodapés: fpage em branco em publicação contínua deixaria o
    # rodapé sem número, e o campo sempre reflete o reinício da numeração.
    styles.apply(paragraph, state.doc, "footer")
    _add_field(paragraph, "PAGE \\* MERGEFORMAT")
    text = vol_issue_year(meta)
    if text:
        paragraph.add_run(f" | {text}")


def _add_field(paragraph, instruction):
    """Campo do Word (ex. PAGE) como run: fldChar begin, instrText, fldChar end."""
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instruction), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.text = text
        run._element.append(element)
    return run


# -----------------
# Cabeçalho com logo
# -----------------


def _logo_header(state, header, spec, meta):
    """
    Cabeçalho da 1ª página com logo (PNG ou SVG):
        left:       [logo | título | DOI]
        right:      [título / DOI | logo]
        full_width: logo na largura do texto e, abaixo, [título | DOI]
    Sem título (show_title=False), só o DOI ao lado do logo. Devolve False,
    sem tocar no documento, quando o arquivo não é PNG/SVG legível.

    Tamanho do logo e larguras das colunas: layout.logo.logo_header (que
    também emite os avisos de redução e de baixa resolução).
    """
    try:
        kind, px_w, px_h = logo_render.read_logo_size(spec.path)
    except ValueError as exc:
        logger.warning("%s; usando o título em texto", exc)
        return False
    fallback = None
    if kind == "svg" and spec.fallback_path:
        try:
            logo_render.read_png_size(spec.fallback_path)
            fallback = spec.fallback_path
        except ValueError as exc:
            logger.warning("%s; usando PNG de reserva transparente", exc)

    # medidas decididas no layout (tamanho do logo, colunas, avisos)
    measures = logo_layout.logo_header(
        spec, kind, px_w, px_h, page_layout.content_width(state.page)
    )
    doc = state.doc

    def add_title(cell):
        """Título do periódico em até duas linhas, no estilo do cabeçalho."""
        _add_journal_title(cell.paragraphs[0], meta["journal_title"], cell.width, doc)

    def add_logo(paragraph, alignment):
        """Logo no tamanho calculado, com o alinhamento dado."""
        paragraph.alignment = alignment
        logo_render.add_logo_run(
            paragraph,
            spec.path,
            measures.width,
            measures.height,
            kind=kind,
            fallback_path=fallback,
        )

    if spec.position == "full_width":
        add_logo(
            header.paragraphs[0] if header.paragraphs else header.add_paragraph(),
            WD_ALIGN_PARAGRAPH.CENTER,
        )
        if spec.show_title:
            title_cell, doi_cell = _header_table(state, header, measures.columns)
            add_title(title_cell)
            _add_doi(
                doi_cell.paragraphs[0], meta["doi"], doc, WD_ALIGN_PARAGRAPH.RIGHT
            )
        else:
            _add_doi(
                header.add_paragraph(), meta["doi"], doc, WD_ALIGN_PARAGRAPH.RIGHT
            )
    elif spec.position == "left":
        if spec.show_title:
            logo_cell, title_cell, doi_cell = _header_table(
                state, header, measures.columns
            )
            add_title(title_cell)
        else:
            logo_cell, doi_cell = _header_table(state, header, measures.columns)
        add_logo(logo_cell.paragraphs[0], WD_ALIGN_PARAGRAPH.LEFT)
        _add_doi(doi_cell.paragraphs[0], meta["doi"], doc, WD_ALIGN_PARAGRAPH.RIGHT)
    else:
        text_cell, logo_cell = _header_table(state, header, measures.columns)
        if spec.show_title:
            add_title(text_cell)
            _add_doi(
                text_cell.add_paragraph(), meta["doi"], doc, WD_ALIGN_PARAGRAPH.LEFT
            )
        else:
            _add_doi(
                text_cell.paragraphs[0], meta["doi"], doc, WD_ALIGN_PARAGRAPH.LEFT
            )
        add_logo(logo_cell.paragraphs[0], WD_ALIGN_PARAGRAPH.RIGHT)
    return True


# -----------------
# Tabelas sem borda de cabeçalho
# -----------------


def _header_table(state, container, widths, v_center=True):
    """
    Tabela 1xN sem borda no cabeçalho/rodapé: cada coluna com largura
    própria, para um texto longo quebrar dentro da sua coluna (tabulação só
    posiciona o início do texto). Sem margem lateral nas células e sem o
    parágrafo vazio que o python-docx cria antes dela.
    """
    widths = [int(w) for w in widths]
    table = container.add_table(rows=1, cols=len(widths), width=sum(widths))
    table.autofit = False
    _remove_placeholder_before(table)
    _zero_side_margins(table)
    cells = table.rows[0].cells
    for column, cell, width in zip(table.columns, cells, widths):
        column.width = cell.width = width
        if v_center:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    return tuple(cells)


def _zero_cell_spacing(container):
    """
    Zera o espaço antes e depois dos parágrafos das células das tabelas do
    container (o cabeçalho da 1ª página).

    Sem isso, os parágrafos herdam o espaço depois do estilo Normal (3,5 mm
    no layout.docx do corpus), e o título do periódico em 3 linhas passava
    do fundo do cabeçalho e empurrava o corpo para baixo.
    """
    for table in container.tables:
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = 0
                paragraph.paragraph_format.space_after = 0


def _remove_placeholder_before(table):
    """Remove o parágrafo vazio que o python-docx deixa antes da tabela (empurraria a tabela para baixo)."""
    previous = table._tbl.getprevious()
    if previous is None or previous.tag != qn("w:p"):
        return
    placeholder = Paragraph(previous, table._parent)
    if not placeholder.text and not placeholder.runs:
        previous.getparent().remove(previous)


def _zero_side_margins(table):
    """Zera as margens laterais das células, para o texto alinhar com o resto da página."""
    margins = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), "0")
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    # ordem do schema CT_TblPr: tblCellMar antes de tblLook
    look = table._tbl.tblPr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(margins)
    else:
        table._tbl.tblPr.append(margins)
