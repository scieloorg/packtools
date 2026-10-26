"""
Render (blocos -> DOCX): figuras, seções, tabelas, texto/fórmulas e a página renderizada.

Migra os testes do old_pdf:
- tests/sps/formats/old_pdf/renderer/docx/test_figure.py
- tests/sps/formats/old_pdf/renderer/docx/test_section.py
- tests/sps/formats/old_pdf/renderer/docx/test_table.py
- tests/sps/formats/old_pdf/renderer/docx/test_text.py
- tests/sps/formats/old_pdf/test_rendered_page_layout.py
"""

import io
import os
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch
from urllib.error import URLError

from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree
from PIL import Image

from packtools.sps.formats.pdf import assets as assets_module
from packtools.sps.formats.pdf.config import page as page_config
from packtools.sps.formats.pdf.assets import EMU_PER_INCH, Assets
from packtools.sps.formats.pdf.layout import page as page_layout
from packtools.sps.formats.pdf.layout.placement import EMPTY_COLUMN_WIDTH, figure_layout, place
from packtools.sps.formats.pdf.render.blocks import _picture, figure, formula, paragraph, render_blocks
from packtools.sps.formats.pdf.render.document import State, new_document
from packtools.sps.formats.pdf.render.inline import add_runs
from packtools.sps.formats.pdf.render.table import FOOT_FONT_SIZE, add_caption, add_table
from packtools.sps.utils import xml_utils

from tests.sps.formats.pdf.helpers import DATA, FIXTURES_DIR, NS, docx, extract, xml

_OMML_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
_W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
_MATH = '<mml:math><mml:mi>x</mml:mi></mml:math>'


def _state(base_layout=DATA['base_layout'], assets_dir=None, page=None):
    """State novo: documento com os estilos do layout base, página padrão e imagens de assets_dir."""
    return State(new_document(base_layout), page or page_config.load_page(), Assets(assets_dir))


def _omml():
    return etree.fromstring(f'<m:oMath xmlns:m="{_OMML_NS}"><m:r><m:rPr/><m:t>x</m:t></m:r></m:oMath>')


def _figure_block(href, **extra):
    return {'type': 'figure', 'id': 'f1', 'href': href, 'alt': '', 'caption': [{'text': 'Figure 1', 'bold': True}],
            'caption_role': 'figure_caption', 'caption_first': False, **extra}


def _placed(block, page=None, layout=None):
    """O bloco com layout e medidas anotados por place (layout: força um layout), como no pipeline."""
    return place([block], page or page_config.load_page(), Assets(), force=layout)[0]


def _sec(body):
    """Blocos de um <sec> com o conteúdo dado."""
    return extract(xml(f'<sec {NS}>{body}</sec>'))


class _WithTmp(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.tmp = tmp.name

    def _png(self, name, px_width, px_height=100, dpi=None):
        path = os.path.join(self.tmp, name)
        image = Image.new('RGB', (px_width, px_height), 'white')
        image.save(path, format='PNG', **({'dpi': (dpi, dpi)} if dpi else {}))
        return path


# -----------------
# Figuras (old: renderer/docx/test_figure.py)
# -----------------


class TestAssetsPath(_WithTmp):
    """Resolução do href da imagem (old _resolve_image_path -> Assets.path)."""

    def test_relative_href_resolves_against_assets_dir(self):
        """old: test_figure.py TestResolveImagePath.test_relative_local_path_resolves_against_assets_dir"""
        path = self._png('fig1.png', 10)
        self.assertEqual(Assets(self.tmp).path('fig1.png'), path)

    def test_empty_href_returns_none(self):
        """old: test_figure.py TestResolveImagePath.test_empty_href_returns_none"""
        for href in ('', None):
            with self.subTest(href=href):
                self.assertIsNone(Assets(self.tmp).path(href))

    def test_absolute_path_returned_unchanged(self):
        """old: test_figure.py TestResolveImagePath.test_absolute_local_path_returned_unchanged"""
        path = self._png('fig1.png', 10)
        self.assertEqual(Assets('/some/other/dir').path(path), path)

    def test_missing_local_file_is_logged_once(self):
        assets = Assets(self.tmp)
        with self.assertLogs(assets_module.logger, level='WARNING') as logs:
            self.assertIsNone(assets.path('fig1.tif'))
            self.assertIsNone(assets.path('fig1.tif'))
            self.assertIsNone(assets.natural_size('fig1.tif'))
        self.assertEqual(len(logs.output), 1)
        self.assertIn('fig1.tif', logs.output[0])
        self.assertIn(os.path.join(self.tmp, 'fig1.tif'), logs.output[0])

    def test_remote_url_is_downloaded_once(self):
        """
        old: test_figure.py TestResolveImagePath.test_remote_url_is_downloaded

        A URL é baixada para um cache local (uma vez por Assets) e o caminho
        devolvido é o do arquivo baixado.
        """
        response = MagicMock()
        response.__enter__.return_value = io.BytesIO(b'fake')
        with patch('packtools.sps.formats.pdf.assets.urlopen', return_value=response) as mock_open:
            assets = Assets()
            first = assets.path('https://example.org/fig1.png')
            second = assets.path('https://example.org/fig1.png')
        self.assertEqual(mock_open.call_count, 1)
        self.assertEqual(mock_open.call_args[0][0].full_url, 'https://example.org/fig1.png')
        self.assertEqual(first, second)
        self.assertTrue(first.endswith('.png'))
        with open(first, 'rb') as f:
            self.assertEqual(f.read(), b'fake')
        shutil.rmtree(assets._cache_dir)

    def test_failed_download_is_logged_and_not_retried(self):
        """O href passa pelo layout e pelo render: uma tentativa e um aviso só."""
        url = 'https://example.org/fig1.png'
        with patch('packtools.sps.formats.pdf.assets.urlopen', side_effect=URLError('timed out')) as mock_open:
            assets = Assets()
            with self.assertLogs(assets_module.logger, level='WARNING') as logs:
                self.assertIsNone(assets.path(url))
                self.assertIsNone(assets.path(url))
                self.assertIsNone(assets.natural_size(url))
        self.assertEqual(mock_open.call_count, 1)
        self.assertEqual(len(logs.output), 1)
        self.assertIn(url, logs.output[0])
        self.assertIn('URLError', logs.output[0])
        shutil.rmtree(assets._cache_dir)

    def test_empty_download_is_logged(self):
        response = MagicMock()
        response.__enter__.return_value = io.BytesIO(b'')
        with patch('packtools.sps.formats.pdf.assets.urlopen', return_value=response):
            assets = Assets()
            with self.assertLogs(assets_module.logger, level='WARNING') as logs:
                self.assertIsNone(assets.path('https://example.org/fig1.png'))
                self.assertIsNone(assets.path('https://example.org/fig1.png'))
        self.assertEqual(len(logs.output), 1)
        self.assertIn('sem conteúdo', logs.output[0])
        shutil.rmtree(assets._cache_dir)


class TestFigureCaption(unittest.TestCase):
    """Legenda da figura no estilo próprio 'SCL Figure Caption' (#1346, PR #1358)."""

    def test_caption_uses_dedicated_style(self):
        """old: test_figure.py TestAddCaption.test_caption_uses_dedicated_style"""
        state = _state()
        figure(state, _placed(_figure_block('missing.png')))
        self.assertEqual(state.doc.paragraphs[-1].style.name, 'SCL Figure Caption')
        self.assertEqual(state.doc.paragraphs[-1].text, 'Figure 1')

    def test_dedicated_style_spacing(self):
        """
        old: test_figure.py TestAddCaption.test_dedicated_style_has_single_line_spacing
        old: test_figure.py TestAddCaption.test_dedicated_style_has_space_after_matching_body_paragraph
        """
        style = new_document(DATA['base_layout']).styles['SCL Figure Caption']
        self.assertEqual(style.paragraph_format.line_spacing, 1.0)
        self.assertEqual(style.paragraph_format.space_after, Pt(5.65))


class TestFigureLayout(_WithTmp):
    """Coluna ou largura toda pela largura natural da imagem (old decide_figure_layout -> figure_layout)."""

    def _layout(self, path, **extra):
        return figure_layout(_figure_block(path, **extra), page_config.load_page(), Assets())

    def test_wide_image_gets_full_width_layout(self):
        """old: test_figure.py TestDecideFigureLayoutUnits.test_wide_image_gets_full_width_layout"""
        # ~18,5 cm a 96 DPI: bem mais larga que uma coluna
        self.assertEqual(self._layout(self._png('wide.png', 698, dpi=96)), page_config.SINGLE_COLUMN)

    def test_narrow_image_stays_within_column(self):
        """old: test_figure.py TestDecideFigureLayoutUnits.test_narrow_image_stays_within_column"""
        self.assertEqual(self._layout(self._png('narrow.png', 100, dpi=96)), page_config.DOUBLE_COLUMN)

    def test_untagged_image_uses_print_resolution_fallback(self):
        """old: test_figure.py TestDecideFigureLayoutUnits.test_untagged_image_uses_print_resolution_fallback"""
        # sem DPI: 300 DPI (impressão) -> ~5,9 cm, cabe na coluna
        self.assertEqual(self._layout(self._png('nodpi.png', 700)), page_config.DOUBLE_COLUMN)

    def test_per_figure_override_replaces_natural_width_decision(self):
        """
        old: test_figure.py TestDecideFigureLayoutUnits.test_layout_dpi_override_replaces_embedded_dpi

        diferença intencional: não há mais 'layout_dpi_override' (chave que
        nenhuma extração preenchia); o ajuste por figura é o override de
        layout por id (data['table_layout_overrides'] -> place(overrides=...)),
        que prevalece sobre o DPI embutido.
        """
        path = self._png('dpi72.png', 612, dpi=72)  # ~21,6 cm a 72 DPI
        page = page_config.load_page()
        self.assertEqual(place([_figure_block(path)], page, Assets())[0]['layout'], page_config.SINGLE_COLUMN)
        block = place([_figure_block(path)], page, Assets(), overrides={'f1': page_config.DOUBLE_COLUMN})[0]
        self.assertEqual(block['layout'], page_config.DOUBLE_COLUMN)


class TestPictureSize(_WithTmp):
    """
    Tamanho da figura: decidido no layout (layout.page.fit_image, anotado por
    place) e só aplicado no render (render.blocks._picture).
    old _natural_width_capped -> layout.page.fit_image.
    """

    def test_natural_width_shrinks_to_ceiling_when_larger(self):
        """old: test_figure.py TestNaturalWidthCapped.test_natural_width_shrinks_to_ceiling_when_larger"""
        self.assertEqual(page_layout.fit_image((Cm(50), Cm(25)), Cm(10), Cm(100)), (Cm(10), Cm(5)))

    def test_natural_width_kept_when_smaller_than_ceiling(self):
        """old: test_figure.py TestNaturalWidthCapped.test_natural_width_kept_when_smaller_than_ceiling"""
        block = _placed(_figure_block(self._png('small.png', 200, dpi=96)))
        # pHYs guarda pixels por metro (inteiro): o DPI relido é aproximado
        self.assertAlmostEqual(block['width'], int(Cm(200 / 96 * 2.54)), delta=5000)
        state = _state()
        figure(state, block)
        self.assertEqual(state.doc.inline_shapes[0].width, block['width'])

    def test_tall_image_shrinks_to_content_height_keeping_proportion(self):
        """Altura acima da útil: as duas medidas reduzidas na mesma proporção (antes, no render)."""
        self.assertEqual(page_layout.fit_image((Cm(10), Cm(40)), Cm(10), Cm(20)), (Cm(5), Cm(20)))
        page = page_config.load_page()
        block = _placed(_figure_block(self._png('tall.png', 100, px_height=2000, dpi=10)), page)
        self.assertEqual(block['height'], page_layout.content_height(page))
        state = _state()
        figure(state, block)
        shape = state.doc.inline_shapes[0]
        self.assertEqual((shape.width, shape.height), (block['width'], block['height']))

    def test_unreadable_image_gets_ceiling_width_and_no_height(self):
        """Sem tamanho natural: largura toda e altura a cargo do python-docx."""
        page = page_config.load_page()
        block = _placed(_figure_block('/no/such/file.png'), page)
        self.assertEqual(block['layout'], page_config.SINGLE_COLUMN)
        self.assertEqual((block['width'], block['height']), (page_layout.content_width(page), None))

    def test_missing_file_prints_alt_text(self):
        """
        old: test_figure.py TestNaturalWidthCapped.test_missing_file_falls_back_to_ceiling

        diferença intencional: sem arquivo não há largura a calcular
        (Assets.natural_width é None) nem imagem a inserir (_picture devolve
        None); a figura imprime o texto alternativo no lugar.
        """
        state = _state()
        self.assertIsNone(Assets().natural_width('/no/such/file.png'))
        self.assertIsNone(_picture(state, '/no/such/file.png', Cm(10), None))
        figure(state, _placed(_figure_block('/no/such/file.png', alt='Texto alternativo')))
        self.assertEqual(len(state.doc.inline_shapes), 0)
        self.assertEqual([p.text for p in state.doc.paragraphs][-2:], ['Texto alternativo', 'Figure 1'])


class TestFigureInsertedPicture(_WithTmp):
    """Imagem inserida pelo handler de figura."""

    def test_inserted_width_uses_assets_dpi_not_python_docx_default(self):
        """
        old: test_figure.py TestAddFigureInsertedWidth.test_inserted_picture_width_matches_infer_image_dpi_not_python_docx_default

        Sem DPI no arquivo vale o fallback de Assets (300 DPI), e não os 72 DPI
        que o python-docx usaria sem largura explícita.
        """
        state = _state()
        figure(state, _placed(_figure_block(self._png('no_dpi.png', 200))))
        self.assertEqual(len(state.doc.inline_shapes), 1)
        width = state.doc.inline_shapes[0].width
        self.assertEqual(width, int(200 / 300 * EMU_PER_INCH))
        self.assertNotEqual(width, int(Cm(200 / 72 * 2.54)))

    def test_inserted_picture_has_zero_wrap_distance(self):
        """old: test_figure.py TestAddFigureInsertedWidth.test_inserted_picture_has_zero_wrap_distance"""
        # sem dist*, o LibreOffice desloca a figura 0,32 cm para a direita (#1395)
        state = _state()
        figure(state, _placed(_figure_block(self._png('fig.png', 200))))
        inline = state.doc.inline_shapes[0]._inline
        for side in ('distT', 'distB', 'distL', 'distR'):
            with self.subTest(side=side):
                self.assertEqual(inline.get(side), '0')


# -----------------
# Seções (old: renderer/docx/test_section.py)
# -----------------


class TestBodySection(unittest.TestCase):
    """Seção do corpo com quebra contínua: capa e corpo dividem a página 1 (#1295/#1296)."""

    def test_body_section_uses_continuous_break(self):
        """old: test_section.py TestGetOrCreateSecondSection.test_second_section_uses_continuous_break"""
        self.assertEqual(_state().body_section().start_type, WD_SECTION.CONTINUOUS)

    def test_reuses_existing_body_section(self):
        """old: test_section.py TestGetOrCreateSecondSection.test_reuses_existing_second_section_without_adding_more"""
        state = _state()
        state.body_section()
        state.body_section()
        self.assertEqual(len(state.doc.sections), 2)

    def test_start_type_is_configurable(self):
        """old: test_section.py TestGetOrCreateSecondSection.test_start_type_is_configurable"""
        page = dict(page_config.load_page(), start_type=WD_SECTION.NEW_PAGE)
        self.assertEqual(_state(page=page).body_section().start_type, WD_SECTION.NEW_PAGE)


class TestRunningHeader(unittest.TestCase):
    """
    Cabeçalho corrente no cabeçalho padrão da 1ª seção, herdado (vinculado)
    pelas seguintes: o LibreOffice não mostra cabeçalho de seção contínua desvinculada.
    """

    @classmethod
    def setUpClass(cls):
        cls.doc = docx(xml_utils.get_xml_tree(str(FIXTURES_DIR / 'a1.xml')))

    def test_running_header_targets_first_section(self):
        """old: test_section.py TestDefaultHeaderInheritance.test_get_default_header_targets_first_section"""
        cells = self.doc.sections[0].header.tables[0].rows[0].cells
        self.assertIn('Acta', cells[0].text)
        self.assertTrue(cells[1].text.strip())

    def test_later_sections_header_stays_linked(self):
        """old: test_section.py TestDefaultHeaderInheritance.test_second_section_header_stays_linked_after_setup"""
        self.assertGreater(len(self.doc.sections), 1)
        for index, section in enumerate(self.doc.sections[1:], 1):
            with self.subTest(section=index):
                self.assertTrue(section.header.is_linked_to_previous)


# -----------------
# Tabelas (old: renderer/docx/test_table.py)
# -----------------


def _table_block(ncols=1, foot=()):
    """Bloco de tabela como sai da extração (sem layout nem medidas: ver _placed)."""
    cells = [{'row': 0, 'col': c, 'rowspan': 1, 'colspan': 1, 'header': True, 'runs': [{'text': f'Col{c}'}]}
             for c in range(ncols)]
    return {'type': 'table', 'id': 't1', 'panel': 0, 'caption': [{'text': 'Table 1', 'bold': True}],
            'nrows': 1, 'ncols': ncols, 'cells': cells, 'foot': [[{'text': t}] for t in foot]}


def _table_width_twips(table):
    return int(table._element.tblPr.find(qn('w:tblW')).get(qn('w:w')))


class TestTableWidth(unittest.TestCase):
    """Largura da tabela respeita o número de colunas da página (não um 2 fixo)."""

    def _page(self, count):
        return dict(page_config.load_page(), column_count=count)

    def test_column_width_by_column_count(self):
        """
        old: test_table.py TestComputeTableWidth.test_two_columns_matches_legacy_formula
        old: test_table.py TestComputeTableWidth.test_one_column_equals_full_content_width
        old: test_table.py TestComputeTableWidth.test_three_columns_uses_generic_formula
        """
        for count in (1, 2, 3):
            with self.subTest(count=count):
                page = self._page(count)
                spacing = page['column_spacing'] * 635
                expected = int((page_layout.content_width(page) - (count - 1) * spacing) / count)
                self.assertEqual(page_layout.column_width(page), expected)
                if count == 1:
                    self.assertEqual(page_layout.column_width(page), page_layout.content_width(page))
                state = _state(page=page)
                add_table(state, _placed(_table_block(), page))
                self.assertEqual(_table_width_twips(state.doc.tables[0]), int(expected / 635))

    def test_single_column_layout_uses_content_width(self):
        """old: test_table.py TestComputeTableWidth.test_single_column_layout_override_ignores_column_count"""
        for count in (1, 2):
            with self.subTest(count=count):
                page = self._page(count)
                state = _state(page=page)
                add_table(state, _placed(_table_block(), page, layout=page_config.SINGLE_COLUMN))
                self.assertEqual(_table_width_twips(state.doc.tables[0]), int(page_layout.content_width(page) / 635))


class TestTableCellImage(_WithTmp):
    """
    #1442: coluna só com imagem ficava com a largura mínima (0,38 cm) e a
    imagem, limitada só à coluna do corpo, saía da célula (codas-38-5-e20250355).
    """

    def _table(self, href):
        cells = [
            {'row': 0, 'col': 0, 'rowspan': 1, 'colspan': 1, 'header': False, 'runs': [{'text': 'Lábio superior'}]},
            {'row': 0, 'col': 1, 'rowspan': 1, 'colspan': 1, 'header': False, 'runs': [{'image': href, 'alt': ''}]},
        ]
        return dict(_table_block(), nrows=1, ncols=2, cells=cells)

    def _render(self, href):
        page = page_config.load_page()
        block = place([self._table(href)], page, Assets(self.tmp))[0]
        state = _state(assets_dir=self.tmp, page=page)
        add_table(state, block)
        extent = state.doc.tables[0]._element.find('.//' + qn('wp:extent'))
        return block['column_widths'], int(extent.get('cx'))

    def test_image_only_column_is_not_empty(self):
        self._png('small.png', 100, dpi=96)  # 75pt
        widths, image = self._render('small.png')
        self.assertGreater(widths[1], EMPTY_COLUMN_WIDTH)
        self.assertEqual(image, Assets(self.tmp).natural_width('small.png'))
        self.assertLessEqual(image, widths[1])

    def test_wide_image_shrinks_to_the_cell(self):
        self._png('wide.png', 800, dpi=96)  # 600pt, mais que a coluna do corpo
        widths, image = self._render('wide.png')
        self.assertEqual(image, widths[1])


class TestTableCaption(unittest.TestCase):
    """Legenda da tabela com entrelinha simples, mesmo sem o estilo no layout."""

    def test_caption_has_single_line_spacing(self):
        """
        old: test_table.py TestAddCaptionParagraph.test_caption_has_single_line_spacing
        old: test_table.py TestAddCaptionParagraph.test_line_spacing_set_even_when_named_style_is_missing
        """
        for base_layout in (DATA['base_layout'], None):
            with self.subTest(base_layout=base_layout):
                state = _state(base_layout=base_layout)
                caption = add_caption(state, [{'text': 'Table 1', 'bold': True}, {'text': '. A caption'}],
                                      'table_caption')
                self.assertEqual(caption.paragraph_format.line_spacing, 1.0)
                expected_style = 'SCL Table Heading' if base_layout else 'Normal'
                self.assertEqual(caption.style.name, expected_style)


class TestTableFoot(unittest.TestCase):
    """Notas de <table-wrap-foot> saem como parágrafos abaixo da tabela."""

    def test_foot_notes_are_rendered_as_paragraphs(self):
        """old: test_table.py TestAddTableFoot.test_foot_notes_are_rendered_as_paragraphs"""
        state = _state()
        add_table(state, _placed(_table_block(foot=['Source: Authors.', '* p < 0.05.'])))
        paragraphs = state.doc.paragraphs
        self.assertEqual([p.text for p in paragraphs][-2:], ['Source: Authors.', '* p < 0.05.'])
        for p in paragraphs[-2:]:
            self.assertTrue(all(r.italic and r.font.size == Pt(FOOT_FONT_SIZE) for r in p.runs))

    def test_no_foot_notes_adds_no_extra_paragraphs(self):
        """old: test_table.py TestAddTableFoot.test_no_foot_notes_adds_no_extra_paragraphs"""
        state = _state()
        before = len(state.doc.paragraphs)
        add_table(state, _placed(_table_block()))
        # só a legenda
        self.assertEqual(len(state.doc.paragraphs) - before, 1)

    def test_foot_from_xml_reaches_document(self):
        """old: test_table.py TestAddTableFoot.test_foot_notes_are_rendered_as_paragraphs (pelo XML)"""
        state = _state()
        render_blocks(state, place(_sec(
            '<table-wrap id="t1"><label>Table 1</label><caption><title>T</title></caption>'
            '<table><tbody><tr><td>Data1</td></tr></tbody></table><table-wrap-foot>'
            '<fn id="fn1"><p>Source: Authors.</p></fn><fn id="fn2"><p>* p &lt; 0.05.</p></fn>'
            '</table-wrap-foot></table-wrap>'), state.page, state.assets))
        texts = [p.text for p in state.doc.paragraphs]
        self.assertIn('Source: Authors.', texts)
        self.assertIn('* p < 0.05.', texts)


_TH = '<th>{}</th>'
_TD = '<td>{}</td>'


def _table_xml(head, body):
    thead = f'<thead><tr>{"".join(_TH.format(h) for h in head)}</tr></thead>' if head is not None else ''
    return (f'<table-wrap id="t1"><label>Table 1</label><caption><title>T</title></caption><table>{thead}'
            f'<tbody><tr>{"".join(_TD.format(c) for c in body)}</tr></tbody></table></table-wrap>')


class TestTableColumnCount(unittest.TestCase):
    """Número de colunas = o maior entre cabeçalho e corpo (#1371)."""

    def test_ncols_is_widest_of_header_and_body(self):
        """
        old: test_table.py TestDetermineNumColsHeaderBodyMismatch.test_body_wider_than_header_uses_body_width
        old: test_table.py TestDetermineNumColsHeaderBodyMismatch.test_header_wider_than_body_uses_header_width
        old: test_table.py TestDetermineNumColsHeaderBodyMismatch.test_no_spans_falls_back_to_widest_of_headers_and_rows
        """
        cases = (
            (['c0', 'c1'], ['c0', 'c1', 'c2'], 3),
            (['c0', 'c1', 'c2', 'c3'], ['c0', 'c1'], 4),
            (None, ['A', 'B', 'C'], 3),  # sem <thead>: só <tr> soltos
        )
        for head, body, expected in cases:
            with self.subTest(head=head, body=body):
                (block,) = [b for b in _sec(_table_xml(head, body)) if b['type'] == 'table']
                self.assertEqual(block['ncols'], expected)


class TestTableHeaderBodyMismatch(unittest.TestCase):
    """Tabela com thead de 2 colunas e tbody de 3 (sem colspan), como no corpus real (#1371)."""

    def _render(self):
        state = _state()
        blocks = place(_sec(_table_xml(['', 'H1'], ['row-label', 'A', 'B'])), state.page, state.assets)
        render_blocks(state, blocks)
        return state.doc

    def test_does_not_raise(self):
        """old: test_table.py TestAddTableHeaderBodyColumnMismatch.test_does_not_raise"""
        self.assertEqual(len(self._render().tables), 1)

    def test_table_has_enough_columns_for_body(self):
        """old: test_table.py TestAddTableHeaderBodyColumnMismatch.test_table_has_enough_columns_for_body"""
        table = self._render().tables[0]
        self.assertEqual(len(table.columns), 3)
        self.assertEqual(table.rows[-1].cells[2].text, 'B')


# -----------------
# Texto e fórmulas (old: renderer/docx/test_text.py)
# -----------------


class TestInlineFormula(unittest.TestCase):
    """Run 'omml' entra como XML bruto no parágrafo, não como run de texto."""

    def test_formula_is_appended_as_raw_xml_not_a_run(self):
        """old: test_text.py TestAddParagraphWithSegmentsFormula.test_formula_segment_is_appended_as_raw_xml_not_a_run"""
        state = _state()
        omml = _omml()
        para = add_runs(state.doc.add_paragraph(), [{'omml': omml}], state)
        self.assertEqual(len(para.runs), 0)
        self.assertIn(omml, list(para._p))

    def test_formula_gets_paragraph_font_size_and_family(self):
        """old: test_text.py TestAddParagraphWithSegmentsFormula.test_formula_segment_gets_paragraph_font_size_and_family"""
        state = _state()
        omml = _omml()
        paragraph(state, {'type': 'paragraph', 'role': 'body', 'runs': [{'omml': omml}], 'indent': 0,
                          'preserve': False})
        self.assertEqual(state.doc.paragraphs[-1].style.name, 'SCL Paragraph')
        w_rpr = omml.find(f'{{{_OMML_NS}}}r').find(f'{{{_W_NS}}}rPr')
        self.assertIsNotNone(w_rpr)
        self.assertIsNotNone(w_rpr.find(f'{{{_W_NS}}}sz'))
        self.assertEqual(w_rpr.find(f'{{{_W_NS}}}rFonts').get(f'{{{_W_NS}}}ascii'), 'Noto Serif')

    def test_text_before_and_after_formula_render_as_runs(self):
        """old: test_text.py TestAddParagraphWithSegmentsFormula.test_text_segments_before_and_after_formula_still_render_as_runs"""
        state = _state()
        omml = _omml()
        para = add_runs(state.doc.add_paragraph(),
                        [{'text': 'Total biomass:'}, {'omml': omml}, {'text': ' (3)'}], state)
        self.assertEqual([r.text for r in para.runs], ['Total biomass:', ' (3)'])
        self.assertIn(omml, list(para._p))


class TestDisplayFormulaParagraphs(unittest.TestCase):
    """
    Fórmula em bloco (<disp-formula>) em parágrafo próprio "SCL Formula",
    separado do texto antes dela; fórmula inline fica no parágrafo do texto (#1385).
    A separação agora é feita na extração (bloco 'formula'); aqui o XML
    passa pela extração e pelo render.
    """

    def _render(self, body, state=None):
        state = state or _state()
        before = len(state.doc.paragraphs)
        render_blocks(state, _sec(body))
        return state.doc.paragraphs[before:]

    def test_display_formula_alone_uses_formula_style(self):
        """old: test_text.py TestAddParagraphWithSegmentsFormulaStyle.test_display_formula_alone_uses_formula_style"""
        state = _state()
        formula(state, {'type': 'formula', 'runs': [{'omml': _omml()}]})
        self.assertEqual(len(state.doc.paragraphs), 1)
        self.assertEqual(state.doc.paragraphs[0].style.name, 'SCL Formula')
        added = self._render(f'<p><disp-formula>{_MATH}</disp-formula></p>')
        self.assertEqual([p.style.name for p in added], ['SCL Formula'])

    def test_text_before_display_formula_goes_to_its_own_paragraph(self):
        """
        old: test_text.py TestAddParagraphWithSegmentsFormulaStyle.test_text_before_display_formula_goes_to_its_own_paragraph

        O "(3)" é o <label> da fórmula (como no SPS) e fica no parágrafo da fórmula.
        """
        added = self._render(f'<p>Total plant biomass: <disp-formula id="e3"><label>(3)</label>{_MATH}'
                             '</disp-formula></p>')
        self.assertEqual([(p.style.name, p.text) for p in added],
                         [('SCL Paragraph', 'Total plant biomass:'), ('SCL Formula', ' (3)')])
        self.assertIsNotNone(added[1]._p.find(f'{{{_OMML_NS}}}oMath'))

    def test_blank_text_before_display_formula_does_not_create_empty_paragraph(self):
        """old: test_text.py TestAddParagraphWithSegmentsFormulaStyle.test_blank_text_before_display_formula_does_not_create_empty_paragraph"""
        self.assertEqual(len(self._render(f'<p>  <disp-formula>{_MATH}</disp-formula></p>')), 1)

    def test_inline_formula_before_display_formula_is_kept(self):
        """old: test_text.py TestAddParagraphWithSegmentsFormulaStyle.test_inline_formula_before_display_formula_is_kept"""
        added = self._render(f'<p><inline-formula>{_MATH}</inline-formula><disp-formula>{_MATH}</disp-formula></p>')
        self.assertEqual([p.style.name for p in added], ['SCL Paragraph', 'SCL Formula'])
        for p in added:
            self.assertIsNotNone(p._p.find(f'{{{_OMML_NS}}}oMath'))

    def test_inline_formula_and_plain_text_keep_default_style(self):
        """
        old: test_text.py TestAddParagraphWithSegmentsFormulaStyle.test_inline_formula_keeps_text_paragraph_and_default_style
        old: test_text.py TestAddParagraphWithSegmentsFormulaStyle.test_paragraph_without_formula_keeps_default_style
        """
        for body in (f'<p>Where <inline-formula>{_MATH}</inline-formula> is the mean value.</p>',
                     '<p>Plain paragraph.</p>'):
            with self.subTest(body=body):
                added = self._render(body)
                self.assertEqual([p.style.name for p in added], ['SCL Paragraph'])

    def test_display_formula_in_reference_keeps_reference_style(self):
        """
        old: test_text.py TestAddParagraphWithSegmentsFormulaStyle.test_explicit_non_default_style_is_not_overridden

        Fórmula em bloco dentro de uma referência fica no mesmo parágrafo, no
        estilo da referência.
        """
        added = self._render(f'<ref-list><ref id="B1"><mixed-citation>x <disp-formula>{_MATH}</disp-formula>'
                             '</mixed-citation></ref></ref-list>')
        self.assertEqual(added[-1].style.name, 'SCL Paragraph Reference')
        self.assertIsNotNone(added[-1]._p.find(f'{{{_OMML_NS}}}oMath'))

    def test_formula_style_missing_from_template_falls_back_to_body_style(self):
        """
        old: test_text.py TestAddParagraphWithSegmentsFormulaStyle.test_falls_back_to_requested_style_when_formula_style_missing_from_template

        diferença intencional: a separação texto/fórmula é feita na extração
        e não depende do layout; sem "SCL Formula" no layout, o parágrafo da
        fórmula usa o estilo do corpo (SCL Paragraph) em vez de juntar texto e
        fórmula num parágrafo só.
        """
        state = _state()
        state.doc.styles['SCL Formula'].delete()
        added = self._render(f'<p>x<disp-formula>{_MATH}</disp-formula></p>', state)
        self.assertNotIn('SCL Formula', state.doc.styles)
        self.assertEqual([(p.style.name, p.text) for p in added], [('SCL Paragraph', 'x'), ('SCL Paragraph', '')])


# -----------------
# Página renderizada (old: test_rendered_page_layout.py)
# -----------------


def _libreoffice():
    return shutil.which('libreoffice') or shutil.which('soffice')


@unittest.skipUnless(_libreoffice() and shutil.which('pdftotext'), 'requer LibreOffice e pdftotext')
class TestRenderedPageLayout(unittest.TestCase):
    """
    O corpo começa na página 1 (quebra contínua) e a página 2 tem a
    numeração editorial (#1295/#1296); sem fpage, numeração natural (#1302).
    Só o PDF renderizado mostra isso: o .docx parece certo nos dois casos.
    """

    def _pages(self, name):
        from packtools.sps.formats.pdf.convert import convert_docx_to_pdf

        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / 'out.docx'
            docx(xml_utils.get_xml_tree(str(FIXTURES_DIR / name))).save(docx_path)
            pdf = convert_docx_to_pdf(str(docx_path), _libreoffice())
            return [subprocess.run(['pdftotext', '-f', str(p), '-l', str(p), '-layout', pdf, '-'],
                                   capture_output=True, text=True, check=True).stdout for p in (1, 2)]

    def test_body_starts_page_one_with_editorial_numbering(self):
        """old: test_rendered_page_layout.py TestBodyStartsPageOneRenderedPageNumbers.test_body_starts_page_one_with_editorial_numbering"""
        page1, page2 = self._pages('a1.xml')
        self.assertIn('271 | VOL.', page1)
        self.assertIn('272 | VOL.', page2)
        self.assertIn('biodiversity', page1.lower())

    def test_body_starts_page_one_without_fpage_numbers_sheets_naturally(self):
        """old: test_rendered_page_layout.py TestBodyStartsPageOneRenderedPageNumbers.test_body_starts_page_one_without_fpage_numbers_sheets_naturally"""
        page1, page2 = self._pages('a4.xml')
        self.assertIn('1 | VOL.', page1)
        self.assertIn('2 | VOL.', page2)


if __name__ == '__main__':
    unittest.main()
