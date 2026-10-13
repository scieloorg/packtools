"""
Migração de tests/sps/formats/old_pdf/pipeline/test_docx.py (todas as classes).

Os pipes do old_pdf (docx_*_pipe, _format_*, _add_two_column_header_table,
_body_column_count, _setup_body_section, _render_tables/_render_figures)
viraram, no formats/pdf, extração em blocos + render. Cada teste antigo é
verificado aqui pelo DOCX gerado (pipeline_docx) ou pelo helper de render
equivalente (render/front.py, render/document.py, config/page.py). A origem de
cada teste está no docstring ("old: Classe.teste").
"""

import json
import tempfile
import unittest
from pathlib import Path

from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from packtools.sps.formats.pdf.config import page as page_config
from packtools.sps.formats.pdf.assets import Assets
from packtools.sps.formats.pdf.render import front as front_render
from packtools.sps.formats.pdf.render.document import State, new_document
from packtools.sps.utils import xml_utils

from .helpers import FIXTURES_DIR, article_xml, docx, paragraphs

# fração da largura do cabeçalho da 1ª página para o título do periódico
# (regra herdada do old_pdf, _JOURNAL_TITLE_DOI_SPLIT); fixada aqui porque a
# constante muda de módulo com a refatoração em curso do render/layout
JOURNAL_TITLE_DOI_SPLIT = 0.65
DOI = '10.1590/example'
DOI_URL = 'http://dx.doi.org/10.1590/example'
TABLE_WRAP = (
    '<table-wrap id="t1"><label>Table 1</label><table><thead><tr><th>a</th><th>b</th></tr></thead>'
    '<tbody><tr><td>1</td><td>2</td></tr></tbody></table></table-wrap>'
)
# anexo: garante uma 3ª seção (página nova) para conferir a numeração em todas
TABLE_BACK = '<app-group><app id="app1"><label>Annex 1</label><p>Texto</p></app></app-group>'
FIG = '<fig id="f1"><label>Figure 1</label><caption><title>Cap</title></caption><graphic xlink:href="missing.jpg"/></fig>'


# -----------------
# Utilitários locais
# -----------------


def _front(volume='10', issue='2', year='2023', fpage='123', lpage='130', elocation='', title='Some Article Title'):
    """Conteúdo de <article-meta> com DOI, título, data, volume/número e local."""
    location = (
        f'<fpage>{fpage}</fpage><lpage>{lpage}</lpage>' if fpage else f'<elocation-id>{elocation}</elocation-id>'
    )
    return (
        f'<article-id pub-id-type="doi">{DOI}</article-id>'
        f'<title-group><article-title>{title}</article-title></title-group>'
        f'<pub-date date-type="collection" publication-format="electronic"><year>{year}</year></pub-date>'
        + (f'<volume>{volume}</volume>' if volume else '')
        + (f'<issue>{issue}</issue>' if issue else '')
        + location
    )


def _article(journal_title='Revista', front=None, body='', back=''):
    """Artigo mínimo (helpers.article_xml) com o título do periódico dado."""
    root = article_xml(_front() if front is None else front, body, back)
    root.find('.//journal-title').text = journal_title
    return root


def _state():
    """Documento vazio (sem os estilos SCL, que só importam nos testes via pipeline) para os helpers de render."""
    return State(new_document(), page_config.load_page(), Assets())


def _meta(**values):
    """Metadados (formato de extract_meta) com valores vazios por padrão."""
    meta = dict(
        journal_title='Journal Title', title_text='Some Article Title', doi='', cite_as='',
        volume='', issue='', year='', fpage='', lpage='', location_label='',
    )
    meta.update(values)
    return meta


def _headers_footers(**values):
    """Documento só com cabeçalhos/rodapés (render_headers_footers) dos metadados dados."""
    state = _state()
    front_render.render_headers_footers(state, _meta(**values))
    return state.doc


def _content_width():
    """Largura útil (EMU) da página padrão: largura menos as margens."""
    page = page_config.load_page()
    return int(page['page_width'] - page['left_margin'] - page['right_margin'])


def _fast_docx(root, **data):
    """DOCX sem os estilos do layout base (bem mais rápido): para testes que só leem texto/seções."""
    return docx(root, base_layout=None, **data)


def _cols_num(section):
    return section._sectPr.find(qn('w:cols')).get(qn('w:num'))


def _start_page(section):
    return section._sectPr.find(qn('w:pgNumType')).get(qn('w:start'))


def _is_borderless_default(table):
    """Sem estilo de tabela nem bordas explícitas: a tabela sai sem borda."""
    tbl_pr = table._tbl.tblPr
    return tbl_pr.find(qn('w:tblStyle')) is None and tbl_pr.find(qn('w:tblBorders')) is None


# Footer data dos testes antigos -> metadados equivalentes
WITH_PAGES = dict(volume='10', issue='2', year='2023', fpage='123', lpage='130', location_label='123-130')
WITH_ELOCATION = dict(volume='33', issue='3', year='2024', location_label='e282794')
WITHOUT_ISSUE = dict(volume='86', issue='', year='2026', location_label='e301043')
WITHOUT_VOLUME = dict(volume='', issue='', year='2023', location_label='e236720')


# -----------------
# Pipeline completo
# -----------------


class TestPipelineDocx(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.a4 = _fast_docx(xml_utils.get_xml_tree(str(FIXTURES_DIR / 'a4.xml')))

    def test_start_page_number(self):
        """
        Numeração começa em fpage; sem fpage, em 1 (em todas as seções).
        O caso com fpage usa um artigo mínimo com fpage 271 (o mesmo do a1.xml,
        que leva ~8 s no pipeline); o caso sem fpage usa o a4.xml.
        old: TestPipelineDocx.test_start_page_number_uses_fpage_when_present
        old: TestPipelineDocx.test_start_page_number_defaults_to_one_without_fpage
        """
        with_fpage = _fast_docx(_article(front=_front(fpage='271', lpage='280'), body='<p>x</p>', back=TABLE_BACK))
        for name, document, expected in (('fpage 271', with_fpage, '271'), ('a4.xml', self.a4, '1')):
            with self.subTest(name):
                self.assertEqual(_start_page(document.sections[0]), expected)
                self.assertEqual({_start_page(s) for s in document.sections}, {expected})

    def test_cite_as_falls_back_to_built_citation_without_editorial_note(self):
        """
        Sem nota "como citar", o rodapé leva a citação montada pelos metadados.
        old: TestPipelineDocx.test_cite_as_falls_back_to_built_citation_without_editorial_note
        """
        text = self.a4.sections[0].first_page_footer.paragraphs[0].text
        self.assertIn('Medeiros CA, Gouveia EM.', text)
        self.assertIn('https://doi.org/10.1590/1982-3533.2024v33n3.282794', text)


# -----------------
# Cabeçalho da 1ª página: título do periódico | DOI
# -----------------


class TestFormatJournalTitleTwoLines(unittest.TestCase):

    def test_journal_title_two_lines(self):
        """
        Primeira palavra numa linha e o resto na segunda (máximo 2 linhas).
        old: TestFormatJournalTitleTwoLines.test_single_word_title_stays_on_one_line
        old: TestFormatJournalTitleTwoLines.test_two_word_title_keeps_one_word_per_line
        old: TestFormatJournalTitleTwoLines.test_four_word_title_is_capped_at_two_lines
        old: TestFormatJournalTitleTwoLines.test_five_word_title_is_capped_at_two_lines
        old: TestFormatJournalTitleTwoLines.test_empty_title_returns_empty_string
        """
        for title, expected in (
            ('Biology', 'Biology'),
            ('Acta Amazonica', 'Acta\nAmazonica'),
            ('Brazilian Journal of Biology', 'Brazilian\nJournal of Biology'),
            ('Urbe. Revista Brasileira de Gestão Urbana', 'Urbe.\nRevista Brasileira de Gestão Urbana'),
            ('', ''),
        ):
            with self.subTest(title=title):
                self.assertEqual(front_render.journal_title_two_lines(title), expected)


class TestFirstPageHeader(unittest.TestCase):
    """Cabeçalho da 1ª página gerado pelo pipeline: tabela [título do periódico | DOI]."""

    @classmethod
    def setUpClass(cls):
        # só o primeiro com os estilos SCL (teste do estilo do run)
        cls.docs = {'Acta Amazonica': docx(_article('Acta Amazonica'))}
        for title in ('Brazilian Journal of Biology', 'Urbe. Revista Brasileira de Gestão Urbana'):
            cls.docs[title] = _fast_docx(_article(title))

    def _header(self, title):
        return self.docs[title].sections[0].first_page_header

    def _cells(self, title):
        return self._header(title).tables[-1].rows[0].cells

    def test_journal_title_in_two_lines_in_left_cell(self):
        """
        O título vai em até 2 linhas no 1º run da célula esquerda da tabela do cabeçalho.
        old: TestJournalTitlePipe.test_two_word_title_keeps_one_word_per_line
        old: TestJournalTitlePipe.test_multi_word_title_is_capped_at_two_lines
        old: TestJournalTitlePipe.test_paragraph_lives_in_the_left_cell_of_a_header_table
        """
        for title, expected in (
            ('Acta Amazonica', 'Acta\nAmazonica'),
            ('Brazilian Journal of Biology', 'Brazilian\nJournal of Biology'),
        ):
            with self.subTest(title=title):
                self.assertEqual(self._cells(title)[0].paragraphs[0].runs[0].text, expected)

    def test_journal_title_run_style(self):
        """old: TestJournalTitlePipe.test_run_uses_the_given_style"""
        run = self._cells('Acta Amazonica')[0].paragraphs[0].runs[0]
        self.assertEqual(run.style.name, 'SCL Journal Title Char')

    def test_left_column_gets_the_configured_share_of_content_width(self):
        """old: TestJournalTitlePipe.test_left_column_gets_the_configured_share_of_content_width"""
        table = self._header('Acta Amazonica').tables[-1]
        expected = int(_content_width() * JOURNAL_TITLE_DOI_SPLIT)
        self.assertAlmostEqual(table.columns[0].width, expected, delta=500)

    def test_doi_in_right_cell_without_tab_and_right_aligned(self):
        """
        DOI como URL na célula direita, sem tabulação antes e alinhado à direita.
        old: TestDocxDoiPipe.test_doi_run_has_no_leading_tab
        old: TestDocxDoiPipe.test_doi_paragraph_is_right_aligned
        """
        para = self._cells('Acta Amazonica')[1].paragraphs[0]
        self.assertEqual(para.runs[-1].text, DOI_URL)
        self.assertEqual(para.alignment, WD_ALIGN_PARAGRAPH.RIGHT)

    def test_journal_title_and_doi_share_one_table(self):
        """old: TestDocxDoiPipe.test_reuses_the_table_journal_title_pipe_already_created"""
        title = 'Urbe. Revista Brasileira de Gestão Urbana'
        header = self._header(title)
        self.assertEqual(len(header.tables), 1)
        cells = header.tables[0].rows[0].cells
        self.assertEqual(cells[0].paragraphs[0].runs[0].text, 'Urbe.\nRevista Brasileira de Gestão Urbana')
        self.assertEqual(cells[1].paragraphs[0].runs[-1].text, DOI_URL)

    def test_doi_without_journal_title(self):
        """
        Sem título do periódico, a tabela do cabeçalho existe mesmo assim (com o DOI).
        old: TestDocxDoiPipe.test_works_standalone_without_journal_title_pipe
        """
        state = _state()
        front_render.render_front(state, _meta(journal_title='', doi=DOI), [])
        header = state.doc.sections[0].first_page_header
        self.assertEqual(len(header.tables), 1)
        self.assertEqual(header.tables[0].rows[0].cells[1].paragraphs[0].runs[-1].text, DOI_URL)

    def test_doi_helper_right_aligns_a_given_paragraph(self):
        """old: TestDocxDoiPipe.test_right_aligns_a_caller_supplied_paragraph_too"""
        state = _state()
        para = state.doc.add_paragraph()
        front_render._add_doi(para, DOI, state.doc, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(para.alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(para.runs[-1].text, DOI_URL)


# -----------------
# Documento com estilos, compartilhado por capa, palavras-chave, corpo e back
# -----------------

COVER_FRONT = (
    '<article-categories><subj-group subj-group-type="heading"><subject>Biologia</subject></subj-group>'
    '</article-categories>'
    '<title-group><article-title>Título</article-title></title-group>'
    '<contrib-group><contrib contrib-type="author"><name><surname>Silva</surname>'
    '<given-names>Ana</given-names></name><xref ref-type="aff" rid="aff1">1</xref></contrib></contrib-group>'
    '<aff id="aff1"><label>1</label><institution content-type="original">Universidade X</institution></aff>'
    '<author-notes><corresp>Correspondência: ana@x.br</corresp></author-notes>'
    '<abstract><title>Resumo</title><p>Texto do resumo.</p></abstract>'
    '<kwd-group xml:lang="pt"><title>Keywords:</title><kwd>one</kwd><kwd>two</kwd><kwd>three</kwd></kwd-group>'
)
BODY = (
    '<sec><title>Intro</title>'
    '<p>A <italic>Genus species</italic> seen<sup>1</sup>.</p>'
    '<p><sup>1</sup> and <sub>2</sub></p>'
    '<p>Plain text.</p></sec>'
)
BACK = (
    '<ack><title>Agradecimentos</title>'
    '<p>ao comitê editorial da <italic>Sociologia &amp; Antropologia</italic>, pela acolhida.</p></ack>'
    '<ref-list><title>Referências</title><ref id="B1">'
    '<mixed-citation>Author A. Title B. Journal C. 2020.</mixed-citation></ref></ref-list>'
)
_SHARED = []


def _shared_docx():
    """DOCX (com estilos) de COVER_FRONT + BODY + BACK, gerado uma vez."""
    if not _SHARED:
        _SHARED.append(docx(_article(front=COVER_FRONT, body=BODY, back=BACK)))
    return _SHARED[0]


def _paragraph(document, text):
    """Primeiro parágrafo do corpo do documento com o texto dado."""
    return next(p for p in document.paragraphs if p.text == text)


# -----------------
# Capa (classes TODO do old_pdf)
# -----------------


class TestCoverBlocks(unittest.TestCase):

    def test_cover_paragraph_styles(self):
        """
        Capa renderizada a partir dos blocos, com os estilos SCL de cada parte.
        old: TestDocxArticleTypeAndCategoryPipe (TODO, sem testes)
        old: TestDocxArticleTitlePipe (TODO, sem testes)
        old: TestDocxAuthorsPipe (TODO, sem testes)
        old: TestDocxAffiliationPipe (TODO, sem testes)
        old: TestDocxCorrespondingPipe (TODO, sem testes)
        old: TestDocxAbstractPipe (TODO, sem testes)
        """
        found = paragraphs(_shared_docx())
        for style, text in (
            ('SCL Article Category', 'Biologia | research-article'),
            ('SCL Article Title', 'Título'),
            ('SCL Abstract Title', 'Resumo'),
            ('SCL Paragraph Abstract', 'Texto do resumo.'),
        ):
            with self.subTest(style=style):
                self.assertIn((style, text), found)
        for style, part in (('SCL Author', 'Silva'), ('SCL Affiliation', 'Universidade X')):
            with self.subTest(style=style):
                self.assertTrue(any(s == style and part in t for s, t in found))
        self.assertTrue(any('ana@x.br' in text for _, text in found))


class TestKeywords(unittest.TestCase):
    """Palavras-chave da capa (#1322)."""

    @classmethod
    def setUpClass(cls):
        cls.para = next(p for p in _shared_docx().paragraphs if p.style.name == 'SCL Paragraph Keywords')

    def test_paragraph_has_space_after(self):
        """old: TestDocxKeyworksPipe.test_paragraph_has_space_after"""
        self.assertGreater(self.para.paragraph_format.space_after.pt, 0)

    def test_title_and_content_runs(self):
        """
        old: TestDocxKeyworksPipe.test_title_and_content_runs
        diferença intencional: cada palavra-chave e cada separador são runs
        próprios (preservam a formatação de cada <kwd>); o texto e os
        estilos de caractere são os mesmos.
        """
        runs = self.para.runs
        self.assertEqual(runs[0].text, 'Keywords: ')
        self.assertEqual(runs[0].style.name, 'SCL Paragraph Keywords Header Char')
        self.assertEqual(''.join(r.text for r in runs[1:]), 'one, two, three')
        self.assertEqual({r.style.name for r in runs[1:]}, {'SCL Paragraph Keywords Char'})


# -----------------
# Rodapé da 1ª página: CITE AS
# -----------------


class TestCiteAsFallback(unittest.TestCase):
    """Sem nota "como citar" nem autores: periódico + volume: local."""

    def _cite(self, **values):
        return _headers_footers(**values).sections[0].first_page_footer.paragraphs[0].text

    def test_cite_as_location(self):
        """
        old: TestDocxCiteAsPipe.test_uses_fpage_lpage_range_when_present
        old: TestDocxCiteAsPipe.test_uses_elocation_id_when_fpage_is_absent
        old: TestFormatCiteAsPartTwo.test_keeps_volume_when_present
        """
        for values, expected in (
            (WITH_PAGES, 'CITE AS: Journal Title 10: 123-130.'),
            (WITH_ELOCATION, 'CITE AS: Journal Title 33: e282794.'),
            (WITHOUT_ISSUE, 'CITE AS: Journal Title 86: e301043.'),
        ):
            with self.subTest(expected=expected):
                text = self._cite(**values)
                self.assertEqual(text, expected)
                self.assertNotIn('-.', text)

    def test_omits_volume_colon_when_volume_is_absent(self):
        """
        Regressão #1349: sem <volume>, nada de ': ' solto.
        old: TestDocxCiteAsPipe.test_omits_volume_colon_when_volume_is_absent
        old: TestFormatCiteAsPartTwo.test_omits_volume_colon_when_volume_is_absent
        """
        text = self._cite(**WITHOUT_VOLUME)
        self.assertEqual(text, 'CITE AS: Journal Title e236720.')
        self.assertNotIn(': e236720', text)

    def test_cite_as_from_xml_without_note(self):
        """
        Mesmo comportamento a partir do XML (sem autores, sem nota).
        old: TestDocxCiteAsPipe.test_uses_elocation_id_when_fpage_is_absent
        old: TestDocxCiteAsPipe.test_omits_volume_colon_when_volume_is_absent
        """
        for front, expected in (
            (_front(volume='33', issue='3', year='2024', fpage='', elocation='e282794'),
             'CITE AS: Journal Title 33: e282794.'),
            (_front(volume='', issue='', fpage='', elocation='e236720'), 'CITE AS: Journal Title e236720.'),
        ):
            with self.subTest(expected=expected):
                document = _fast_docx(_article('Journal Title', front=front))
                self.assertEqual(document.sections[0].first_page_footer.paragraphs[0].text, expected)


class TestCiteAsWithNote(unittest.TestCase):
    """A nota "como citar" já é a citação completa (#1349, 2ª rodada)."""

    NOTE = (
        'Colorado Z., G. J.; Valencia-C, G. 2024. Avifauna of a white-sand '
        'forest in the Colombian Amazon. Acta Amazonica 55: e55bc24197.'
    )

    def test_note_is_used_verbatim_without_appending_journal_or_location(self):
        """old: TestDocxCiteAsPipeWithNote.test_note_is_used_verbatim_without_appending_journal_or_location"""
        document = _headers_footers(
            journal_title='Acta Amazonica', cite_as=self.NOTE, volume='55', year='2024', location_label='e55bc24197'
        )
        text = document.sections[0].first_page_footer.paragraphs[0].text
        self.assertEqual(text, 'CITE AS: ' + self.NOTE)
        self.assertEqual(text.count('e55bc24197'), 1)
        self.assertEqual(text.count('Acta Amazonica'), 1)

    def test_note_from_xml_is_used_verbatim(self):
        """
        Mesma regra a partir do XML: a nota do <fn-group> vai para o rodapé e sai do corpo.
        old: TestDocxCiteAsPipeWithNote.test_note_is_used_verbatim_without_appending_journal_or_location
        """
        back = (
            '<fn-group><fn fn-type="other"><label>How to cite this article</label>'
            f'<p>{self.NOTE}</p></fn></fn-group>'
        )
        front = _front(volume='55', issue='', year='2024', fpage='', elocation='e55bc24197')
        document = _fast_docx(_article('Acta Amazonica', front=front, body='<p>Texto.</p>', back=back))
        text = document.sections[0].first_page_footer.paragraphs[0].text
        self.assertEqual(text, 'CITE AS: ' + self.NOTE)
        self.assertNotIn(self.NOTE, [t for _, t in paragraphs(document)])

    def test_note_without_trailing_punctuation_gets_a_period(self):
        """old: TestDocxCiteAsPipeWithNote.test_note_without_trailing_punctuation_gets_a_period"""
        note = (
            'Sousa VR, Gomes MM, Couri MS (2024) On Cerodontha. Zoologia 41: '
            'e23038. https://doi.org/10.1590/S1984-4689.v41.e23038'
        )
        document = _headers_footers(journal_title='Zoologia', cite_as=note, year='2024')
        text = document.sections[0].first_page_footer.paragraphs[0].text
        self.assertTrue(text.endswith('e23038.'))
        self.assertEqual(text.count('Zoologia'), 1)


# -----------------
# Cabeçalho das demais páginas: periódico | título do artigo
# -----------------


class TestRunningHeader(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        document = docx(_article('Brazilian Journal of Biology'))
        cls.table = document.sections[0].header.tables[-1]
        cls.journal_para, cls.title_para = (c.paragraphs[0] for c in cls.table.rows[0].cells)

    def test_journal_title_is_not_split_into_lines(self):
        """old: TestDocxSecondHeaderPipe.test_journal_title_is_not_split_into_lines"""
        self.assertEqual(self.journal_para.runs[0].text, 'Brazilian Journal of Biology')

    def test_journal_title_uses_the_small_header_style_not_the_masthead_style(self):
        """old: TestDocxSecondHeaderPipe.test_journal_title_uses_the_small_header_style_not_the_masthead_style"""
        self.assertEqual(self.journal_para.runs[0].style.name, 'SCL Header Paragraph Char')

    def test_article_title_is_in_its_own_cell_right_aligned(self):
        """
        old: TestDocxSecondHeaderPipe.test_article_title_is_in_its_own_cell
        old: TestDocxSecondHeaderPipe.test_article_title_is_right_aligned
        """
        self.assertEqual(self.title_para.runs[0].text, 'Some Article Title')
        self.assertEqual(self.title_para.alignment, WD_ALIGN_PARAGRAPH.RIGHT)

    def test_columns_split_content_width_evenly(self):
        """old: TestDocxSecondHeaderPipe.test_columns_split_content_width_evenly"""
        self.assertEqual(self.table.columns[0].width, self.table.columns[1].width)

    def test_table_is_borderless(self):
        """
        old: TestDocxSecondHeaderPipe.test_table_keeps_the_borderless_default_style
        diferença intencional: o novo documento copia os estilos do
        layout.docx e o python-docx não reporta 'Normal Table'; o que importa
        (tabela sem estilo nem bordas próprias) é verificado no XML.
        """
        self.assertTrue(_is_borderless_default(self.table))


# -----------------
# Rodapés: "PÁGINA | VOL. v (n) ano: local"
# -----------------


class TestPageFooters(unittest.TestCase):

    CASES = (
        (WITH_PAGES, 'VOL. 10 (2) 2023: 123-130'),
        (WITH_ELOCATION, 'VOL. 33 (3) 2024: e282794'),
        (WITHOUT_ISSUE, 'VOL. 86 2026: e301043'),
    )

    def _check(self, footer_of):
        for values, expected in self.CASES:
            with self.subTest(expected=expected):
                text = footer_of(_headers_footers(**values)).text
                self.assertEqual(text, f' | {expected}')
                self.assertNotIn(': -', text)
                self.assertNotIn('()', text)

    def test_body_footer(self):
        """
        Rodapé das demais páginas (seção do corpo), depois do campo PAGE.
        old: TestDocxSecondFooterPipe.test_uses_fpage_lpage_range_when_present
        old: TestDocxSecondFooterPipe.test_uses_elocation_id_when_fpage_is_absent
        old: TestDocxSecondFooterPipe.test_omits_issue_parentheses_when_issue_is_absent
        """
        self._check(lambda doc: doc.sections[1].footer.paragraphs[0])

    def test_first_page_footer(self):
        """
        Última linha do rodapé da 1ª página.
        old: TestDocxPageVolIssueYearPipe.test_uses_fpage_lpage_range_when_present
        old: TestDocxPageVolIssueYearPipe.test_uses_elocation_id_when_fpage_is_absent
        old: TestDocxPageVolIssueYearPipe.test_omits_issue_parentheses_when_issue_is_absent
        """
        self._check(lambda doc: doc.sections[0].first_page_footer.paragraphs[-1])

    def test_footers_have_page_field(self):
        """Os dois rodapés de página têm o campo PAGE (complemento dos testes acima)."""
        document = _headers_footers(**WITH_PAGES)
        for paragraph in (document.sections[1].footer.paragraphs[0],
                          document.sections[0].first_page_footer.paragraphs[-1]):
            self.assertIn('PAGE', ''.join(e.text for e in paragraph._p.iter(qn('w:instrText'))))


class TestFormatVolIssueYear(unittest.TestCase):

    def test_vol_issue_year(self):
        """
        old: TestFormatVolIssueYear.test_keeps_both_when_present
        old: TestFormatVolIssueYear.test_omits_issue_parentheses_when_issue_is_absent
        old: TestFormatVolIssueYear.test_omits_vol_label_when_volume_is_absent
        """
        for footer, expected in (
            ({'volume': '10', 'issue': '2', 'year': '2023', 'location_label': '123-130'}, 'VOL. 10 (2) 2023: 123-130'),
            ({'volume': '86', 'issue': '', 'year': '2026', 'location_label': 'e301043'}, 'VOL. 86 2026: e301043'),
            ({'volume': '', 'issue': '67', 'year': '2023', 'location_label': 'e236720'}, '(67) 2023: e236720'),
        ):
            with self.subTest(expected=expected):
                self.assertEqual(front_render.vol_issue_year(footer), expected)


# -----------------
# Corpo
# -----------------


class TestBodyRuns(unittest.TestCase):
    """Cada trecho com formatação vira um run com a mesma formatação (item 04 do backlog)."""

    @classmethod
    def setUpClass(cls):
        cls.paras = {p.text: p for p in _shared_docx().paragraphs if p.style is not None and p.style.name == 'SCL Paragraph'}

    def test_paragraph_emits_one_run_per_segment_with_matching_style(self):
        """old: TestDocxBodyPipe.test_paragraph_emits_one_run_per_segment_with_matching_style"""
        runs = self.paras['A Genus species seen1.'].runs
        self.assertEqual([r.text for r in runs], ['A ', 'Genus species', ' seen', '1', '.'])
        self.assertIsNone(runs[0].italic)
        self.assertTrue(runs[1].italic)
        self.assertIsNone(runs[2].italic)
        self.assertTrue(runs[3].font.superscript)
        self.assertIsNone(runs[4].font.superscript)

    def test_superscript_and_subscript_segments_in_the_same_paragraph(self):
        """old: TestDocxBodyPipe.test_superscript_and_subscript_segments_in_the_same_paragraph"""
        runs = self.paras['1 and 2'].runs
        self.assertTrue(runs[0].font.superscript)
        self.assertTrue(runs[2].font.subscript)

    def test_unstyled_segment_does_not_force_run_formatting(self):
        """old: TestDocxBodyPipe.test_unstyled_segment_does_not_force_run_bold_false"""
        run = self.paras['Plain text.'].runs[0]
        for attr in (run.bold, run.italic, run.font.superscript, run.font.subscript):
            self.assertIsNone(attr)


class TestBackMatter(unittest.TestCase):

    def _para(self, text):
        return _paragraph(_shared_docx(), text)

    def test_reference_has_single_line_spacing(self):
        """old: TestDocxReferencesPipe.test_reference_has_single_line_spacing"""
        para = self._para('Author A. Title B. Journal C. 2020.')
        self.assertEqual(para.style.name, 'SCL Paragraph Reference')
        self.assertEqual(para.paragraph_format.line_spacing, 1.0)

    def test_acknowledgments_title_is_a_section_title(self):
        """old: TestDocxAcknowledgmentsPipe (1ª definição, TODO sem testes; sombreada pela 2ª)"""
        self.assertEqual(self._para('Agradecimentos').style.name, 'SCL Section Title')

    def test_renders_whole_paragraph_with_inline_italic(self):
        """
        #1374: o texto depois do <italic> não se perde.
        old: TestDocxAcknowledgmentsPipe.test_renders_whole_paragraph_with_inline_italic (2ª definição)
        """
        para = self._para('ao comitê editorial da Sociologia & Antropologia, pela acolhida.')
        self.assertEqual([r.text for r in para.runs if r.italic], ['Sociologia & Antropologia'])


# -----------------
# Anexos e material suplementar
# -----------------


class TestAnnex(unittest.TestCase):

    def test_annex_footer_has_no_leading_pipe(self):
        """
        Rodapé dos anexos: só volume/ano/local, sem número de página nem " | ".
        old: TestDocxSupplementaryMaterialPipe.test_footer_has_no_leading_pipe
        old: TestDocxSupplementaryMaterialPipe.test_uses_elocation_id_when_fpage_is_absent
        old: TestDocxSupplementaryMaterialPipe.test_omits_issue_parentheses_when_issue_is_absent
        """
        for values, expected in (
            (dict(volume='53', issue='4', year='2023', location_label='271-280'), 'VOL. 53 (4) 2023: 271-280'),
            (WITH_ELOCATION, 'VOL. 33 (3) 2024: e282794'),
            (WITHOUT_ISSUE, 'VOL. 86 2026: e301043'),
        ):
            with self.subTest(expected=expected):
                state = _state()
                section = state.doc.add_section()
                front_render.annex_footer(state, section, _meta(**values))
                self.assertEqual(section.footer.paragraphs[-1].text, expected)

    @classmethod
    def setUpClass(cls):
        back = (
            '<app-group><app id="app1"><label>Annex 1</label><p>Texto do anexo</p>'
            f'<sec><title>Zone 1</title>{TABLE_WRAP}</sec></app></app-group>'
            '<sec sec-type="supplementary-material"><title>Supplementary Material</title>'
            '<supplementary-material id="s1"><label>Suppl. 1</label>'
            '<media xlink:href="a.pdf" mimetype="application" mime-subtype="pdf"/></supplementary-material></sec>'
        )
        front = _front(volume='53', issue='4', year='2023', fpage='271', lpage='280')
        cls.document = docx(_article(front=front, body='<p>Texto.</p>', back=back))

    def test_annex_footer_from_xml(self):
        """
        Pelo pipeline: seção de anexos em página nova, coluna única, com o rodapé próprio.
        old: TestDocxSupplementaryMaterialPipe.test_footer_has_no_leading_pipe
        """
        section = self.document.sections[-1]
        self.assertEqual(section.start_type, WD_SECTION.NEW_PAGE)
        self.assertEqual(_cols_num(section), '1')
        self.assertEqual(section.footer.paragraphs[-1].text, 'VOL. 53 (4) 2023: 271-280')

    def test_renders_headings_paragraphs_and_tables(self):
        """
        old: TestDocxSupplementaryMaterialPipe.test_renders_headings_paragraphs_and_tables
        diferença intencional: a legenda "Table 1" sai na ordem do documento
        (o teste antigo a ignorava); conferida aqui explicitamente.
        """
        document = self.document
        found = paragraphs(document)
        # parágrafos depois da última quebra de seção (início dos anexos)
        annex = found[max(i for i, (style, _) in enumerate(found) if style is None) + 1:]
        self.assertEqual(annex, [
            ('SCL Section Title', 'Annex 1'),
            ('SCL Paragraph', 'Texto do anexo'),
            ('SCL Subsection Title', 'Zone 1'),
            ('SCL Table Heading', 'Table 1'),
            ('SCL Section Title', 'Supplementary Material'),
            ('SCL Paragraph', 'Suppl. 1: a.pdf (application/pdf)'),
        ])
        self.assertEqual(len(document.tables), 1)


# -----------------
# Tabela de 2 colunas sem borda (cabeçalhos)
# -----------------


class TestHeaderTable(unittest.TestCase):

    def setUp(self):
        self.state = _state()
        self.header = self.state.doc.sections[0].first_page_header

    def _table(self, ratio=0.5):
        width = _content_width()
        front_render._header_table(self.state, self.header, [int(width * ratio), width - int(width * ratio)])
        return self.header.tables[-1]

    def test_default_split_is_even(self):
        """old: TestAddTwoColumnHeaderTable.test_default_split_is_even"""
        table = self._table()
        self.assertEqual(table.columns[0].width, table.columns[1].width)

    def test_custom_ratio_splits_unevenly(self):
        """old: TestAddTwoColumnHeaderTable.test_custom_ratio_splits_unevenly"""
        table = self._table(0.65)
        width = _content_width()
        self.assertAlmostEqual(table.columns[0].width, int(width * 0.65), delta=500)
        self.assertAlmostEqual(table.columns[1].width, width - int(width * 0.65), delta=500)

    def test_table_is_borderless(self):
        """
        old: TestAddTwoColumnHeaderTable.test_table_keeps_the_borderless_default_style
        diferença intencional: verificado como "sem estilo nem bordas" (ver TestRunningHeader).
        """
        self.assertTrue(_is_borderless_default(self._table()))

    def test_removes_the_auto_created_empty_placeholder_paragraph(self):
        """old: TestAddTwoColumnHeaderTable.test_removes_the_auto_created_empty_placeholder_paragraph"""
        self._table()
        self.assertEqual(len(self.header.paragraphs), 0)

    def test_keeps_a_placeholder_paragraph_that_already_has_text(self):
        """old: TestAddTwoColumnHeaderTable.test_keeps_a_placeholder_paragraph_that_already_has_text"""
        self.header.paragraphs[0].add_run('not actually empty')
        self._table()
        self.assertEqual([p.text for p in self.header.paragraphs], ['not actually empty'])

    def test_zeroes_left_and_right_cell_margins(self):
        """old: TestAddTwoColumnHeaderTable.test_zeroes_left_and_right_cell_margins"""
        margins = self._table()._tbl.tblPr.find(qn('w:tblCellMar'))
        self.assertIsNotNone(margins)
        for side in ('left', 'right'):
            with self.subTest(side=side):
                self.assertEqual(margins.find(qn(f'w:{side}')).get(qn('w:w')), '0')


# -----------------
# Colunas do corpo
# -----------------


class TestBodyColumnConfiguration(unittest.TestCase):
    """
    O número de colunas do corpo (e o restaurado depois de tabela/figura de
    largura total) vem de default_column_count do layout de página.
    """

    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        layout = json.loads(page_config.DEFAULT_LAYOUT_PATH.read_text())
        cls.layouts = {}
        cls.docs = {}
        for count in (1, 2, 3):
            layout['page']['default_column_count'] = count
            path = Path(cls._tmp.name) / f'cols{count}.json'
            path.write_text(json.dumps(layout))
            cls.layouts[count] = str(path)

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def _docx(self, count, wide=''):
        """Documento (em cache) com o corpo em `count` colunas; `wide` é forçado à largura total."""
        key = (count, wide)
        if key not in self.docs:
            body = f'<sec><title>S</title><p>antes</p>{wide}<p>depois</p></sec>'
            self.docs[key] = _fast_docx(
                _article(body=body),
                page_layout=self.layouts[count],
                table_layout_overrides={'t1': page_config.SINGLE_COLUMN, 'f1': page_config.SINGLE_COLUMN},
            )
        return self.docs[key]

    def test_body_column_count_defaults_to_two(self):
        """old: TestBodyColumnConfiguration.test_body_column_count_defaults_to_two"""
        self.assertEqual(page_config.load_page()['column_count'], 2)
        self.assertEqual(_cols_num(_fast_docx(_article(body='<p>x</p>')).sections[1]), '2')

    def test_body_column_count_reads_config(self):
        """old: TestBodyColumnConfiguration.test_body_column_count_reads_config"""
        self.assertEqual(page_config.load_page(self.layouts[1])['column_count'], 1)

    def test_body_section_uses_configured_column_count(self):
        """
        old: TestBodyColumnConfiguration.test_setup_body_section_uses_configured_column_count
        old: TestBodyColumnConfiguration.test_setup_body_section_supports_three_columns
        """
        for count in (1, 3):
            with self.subTest(count=count):
                self.assertEqual(_cols_num(self._docx(count).sections[1]), str(count))

    def test_restored_body_section_uses_configured_column_count(self):
        """
        Depois da tabela larga, o corpo volta ao número configurado de colunas.
        old: TestBodyColumnConfiguration.test_restore_body_column_section_restores_configured_column_count
        old: TestBodyColumnConfiguration.test_restore_body_column_section_supports_three_columns
        """
        for count in (1, 3):
            with self.subTest(count=count):
                document = self._docx(count, TABLE_WRAP)
                self.assertEqual(_cols_num(document.sections[-1]), str(count))
                last_text = [t for _, t in paragraphs(document) if t][-1]
                self.assertEqual(last_text, 'depois')

    def test_wide_element_section_switch(self):
        """
        Corpo em 1 coluna: tabela/figura larga não abre seção; em 2 colunas,
        abre uma seção de 1 coluna e outra que restaura as 2 (+2 seções).
        old: TestBodyColumnConfiguration.test_render_tables_skips_section_switch_when_body_is_single_column
        old: TestBodyColumnConfiguration.test_render_tables_switches_section_when_body_has_multiple_columns
        old: TestBodyColumnConfiguration.test_render_figures_skips_section_switch_when_body_is_single_column
        old: TestBodyColumnConfiguration.test_render_figures_switches_section_when_body_has_multiple_columns
        """
        for count, extra in ((1, 0), (2, 2)):
            before = len(self._docx(count).sections)
            for name, wide in (('table', TABLE_WRAP), ('figure', FIG)):
                with self.subTest(count=count, element=name):
                    sections = self._docx(count, wide).sections
                    self.assertEqual(len(sections), before + extra)
                    if extra:
                        self.assertEqual([_cols_num(s) for s in sections[1:]], ['2', '1', '2'])


if __name__ == '__main__':
    unittest.main()
